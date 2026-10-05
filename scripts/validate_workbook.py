#!/usr/bin/env python3
"""Validate a generated ``workbook.xlsx`` against the Workbook Contract (`src/publisher/contracts/workbook_contract.py`).

Checks (CLAUDE.md change-loop step 6): exactly four visible sheets in contract order; detailed-sheet headers
equal the declared contract columns in order; row grains match the package JSON; no formulas in any cell;
financial cells numeric (or an explicit unavailable sentinel), never strings such as ``Save USD ...``; identifier
columns stored as text.

Usage:  python scripts/validate_workbook.py <output_dir>      # expects workbook.xlsx and packages/*.json inside
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any, List, Optional, Sequence

ROOT = Path(__file__).resolve().parents[1]
for _path in (ROOT, ROOT / "src"):
    if str(_path) not in sys.path:
        sys.path.insert(0, str(_path))

from publisher.contracts import workbook_contract as wc  # noqa: E402

DETAILED = (
    (wc.SHEET_PORTFOLIO, wc.PORTFOLIO_COLUMNS, "portfolio"),
    (wc.SHEET_SCENARIOS, wc.SCENARIO_COLUMNS, "scenarios"),
    (wc.SHEET_REGISTRY, wc.REGISTRY_COLUMNS, "candidates"),
)
FIRST_DATA_ROW = 5
HEADER_ROW = 4
SENTINELS = {wc.DISPLAY_NA, wc.DISPLAY_UNAVAILABLE, "MIXED"}


def expected_grains(packages: Sequence[Any]) -> dict:
    return {
        "portfolio": len(packages),
        "scenarios": sum(len(p["scenarios"]) for p in packages),
        "candidates": sum(len(p["candidates"]) for p in packages),
    }


def validate_workbook(path: Path, packages: Optional[Sequence[Any]] = None) -> List[str]:
    from openpyxl import load_workbook

    errors: List[str] = []
    wb = load_workbook(path)  # formulas are preserved as strings (data_only=False), so they are detectable
    names = [ws.title for ws in wb.worksheets]
    if tuple(names) != wc.SHEET_NAMES:
        errors.append(f"sheets {names} != {list(wc.SHEET_NAMES)}")
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            errors.append(f"sheet {ws.title!r} is {ws.sheet_state}, must be visible")
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.data_type != "s" and cell.value.startswith("=")):
                    errors.append(f"{ws.title}!{cell.coordinate}: cell formulas are not allowed")
    grains = expected_grains(packages) if packages is not None else None
    for sheet_name, columns, grain in DETAILED:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = [ws.cell(row=HEADER_ROW, column=i).value for i in range(1, len(columns) + 1)]
        if headers != [c.header for c in columns]:
            errors.append(f"{sheet_name}: header row differs from the declared contract columns")
        if ws.cell(row=HEADER_ROW, column=len(columns) + 1).value is not None:
            errors.append(f"{sheet_name}: unexpected column beyond the {len(columns)} declared columns")
        rows = [r for r in range(FIRST_DATA_ROW, ws.max_row + 1) if ws.cell(row=r, column=1).value is not None]
        if grains is not None and len(rows) != grains[grain]:
            errors.append(f"{sheet_name}: {len(rows)} rows, expected {grains[grain]}")
        for idx, spec in enumerate(columns, start=1):
            for r in rows:
                cell = ws.cell(row=r, column=idx)
                if spec.number_format == "@":
                    if cell.value is not None and cell.data_type != "s":
                        errors.append(f"{sheet_name}!{cell.coordinate}: identifier column {spec.key} must be text")
                elif spec.number_format in (wc.MONEY_FORMAT, wc.PERCENT_FORMAT):
                    v = cell.value
                    numeric = isinstance(v, (int, float)) and not isinstance(v, bool)
                    if not numeric and v not in SENTINELS:
                        errors.append(f"{sheet_name}!{cell.coordinate}: {spec.key} must be numeric or an unavailable sentinel, got {v!r}")
    return errors


def main(argv: Sequence[str]) -> int:
    if len(argv) != 2:
        print(__doc__)
        return 2
    out = Path(argv[1])
    workbook = out / "workbook.xlsx"
    package_files = sorted((out / "packages").glob("*.json"))
    packages = [json.loads(f.read_text()) for f in package_files] if package_files else None
    errors = validate_workbook(workbook, packages)
    if errors:
        print("\n".join(errors[:50]))
        print(f"workbook validation FAILED: {len(errors)} problem(s)")
        return 1
    print(f"workbook validation OK: {workbook} ({len(package_files)} packages)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
