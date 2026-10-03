"""Excel workbook reporter for Workbook Contract v1.2.0.

The workbook is a presentation-only projection of locked Reporter JSON Contract
v1.1.0.  Exact detailed-sheet columns and section groups are owned by
``contracts.workbook_contract``; this renderer cannot silently drop fields.

Rendered with ``openpyxl`` (pip-installable); no proprietary runtime is required.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterator, Mapping, Optional, Sequence

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from contracts import reporter_contract as rc
from contracts import workbook_contract as wc

ESPRESSO = "#3B2F2F"
ESPRESSO_2 = "#5A463B"
CREAM = "#F7F1E8"
CREAM_2 = "#EEE3D3"
GOLD = "#B58A3A"
GREEN = "#DDEBDD"
AMBER = "#FFF0CC"
RED = "#F8D7DA"
BLUE = "#DDEAF6"
GRAY = "#E9E7E4"
WHITE = "#FFFFFF"
TEXT = "#2B2522"
MUTED = "#6B625D"
MONEY_FMT = wc.MONEY_FORMAT
PCT_FMT = wc.PERCENT_FORMAT
INT_FMT = '#,##0'
DECIMAL_FMT = '0.00'

# Excel rejects cells longer than this; verbose audit-JSON cells are cut with a visible marker rather than
# producing a file Excel reports as corrupt. The full content always remains in the package JSON.
EXCEL_MAX_CELL_CHARS = 32767
TRUNCATION_MARKER = " …[truncated; full content in the package JSON]"


def _col_letter(n: int) -> str:
    return get_column_letter(n)


def _rgb(color: str) -> str:
    return color.lstrip("#")


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=_rgb(color), end_color=_rgb(color))


def _font(color: str = TEXT, size: int = 9, bold: bool = False, italic: bool = False) -> Font:
    return Font(color=_rgb(color), size=size, bold=bold, italic=italic)


def _cells(sheet: Any, ref: str) -> Iterator[Any]:
    block = sheet[ref]
    if not isinstance(block, tuple):
        yield block
        return
    for row in block:
        if isinstance(row, tuple):
            yield from row
        else:
            yield row


def _style(
    sheet: Any, ref: str, *, fill: Optional[str] = None, font: Optional[Font] = None, number_format: Optional[str] = None,
    horizontal: Optional[str] = None, vertical: Optional[str] = None, wrap: Optional[bool] = None,
) -> None:
    for cell in _cells(sheet, ref):
        if fill:
            cell.fill = _fill(fill)
        if font:
            cell.font = font
        if number_format:
            cell.number_format = number_format
        if horizontal or vertical or wrap is not None:
            cell.alignment = Alignment(
                horizontal=horizontal or cell.alignment.horizontal, vertical=vertical or cell.alignment.vertical,
                wrap_text=cell.alignment.wrap_text if wrap is None else wrap,
            )


def _row_height(sheet: Any, first_row: int, last_row: int, height: float) -> None:
    for r in range(first_row, last_row + 1):
        sheet.row_dimensions[r].height = height


def _normalize(v: Any) -> Any:
    if v is None:
        return wc.DISPLAY_NA
    if hasattr(v, "value"):
        return v.value
    return v


def _excel_text(v: Any) -> Any:
    """Return the exact visible value for a contract text cell.

    Long digit-only identifiers (workspace IDs) stay text because the cell is written as an explicit string with
    the Text number format; no apostrophe workaround is needed or wanted with this renderer.
    """
    return _normalize(v)


def _put(sheet: Any, row: int, col: int, value: Any) -> Any:
    """Write one cell, guaranteeing that text is never interpreted as a formula."""
    value = _normalize(value)
    if isinstance(value, str) and len(value) > EXCEL_MAX_CELL_CHARS:
        value = value[: EXCEL_MAX_CELL_CHARS - len(TRUNCATION_MARKER)] + TRUNCATION_MARKER
    cell = sheet.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # presentation-only: workbook cells never carry formulas
    return cell


def _add_table(sheet: Any, ref: str, name: str, style: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def _sheet_title(sheet: Any, title: str, subtitle: str, end_col: int) -> None:
    end = _col_letter(max(end_col, 8))
    sheet.merge_cells(f"A1:{end}1")
    sheet["A1"] = title
    _style(sheet, "A1", fill=ESPRESSO, font=_font(WHITE, 16, bold=True), vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells(f"A2:{end}2")
    sheet["A2"] = subtitle
    _style(sheet, "A2", fill=CREAM, font=_font(ESPRESSO_2, 9, italic=True), vertical="center", wrap=True)
    sheet.row_dimensions[2].height = 24


def _write_group_row(sheet: Any, columns: Sequence[wc.ColumnSpec], groups: Sequence[wc.ColumnGroup], row: int = 3) -> None:
    by_key = {c.key: idx + 1 for idx, c in enumerate(columns)}
    for i, group in enumerate(groups):
        start, end = by_key[group.start_key], by_key[group.end_key]
        start_letter, end_letter = _col_letter(start), _col_letter(end)
        sheet[f"{start_letter}{row}"] = group.title
        _style(
            sheet, f"{start_letter}{row}", fill=GOLD if i % 2 == 0 else ESPRESSO_2, font=_font(WHITE, 9, bold=True),
            horizontal="center", vertical="center",
        )
        if start != end:
            sheet.merge_cells(f"{start_letter}{row}:{end_letter}{row}")
    sheet.row_dimensions[row].height = 20


def _write_contract_table(
    sheet: Any,
    *,
    columns: Sequence[wc.ColumnSpec],
    groups: Sequence[wc.ColumnGroup],
    rows: Sequence[Mapping[str, Any]],
    table_name: str,
    freeze_columns: int,
    wrap_keys: Sequence[str] = (),
    json_keys: Sequence[str] = (),
    row_height: float = 24,
) -> None:
    _write_group_row(sheet, columns, groups, 3)
    end_row = 4 + len(rows)
    end_col = len(columns)
    end_letter = _col_letter(end_col)

    for idx, spec in enumerate(columns, start=1):
        _put(sheet, 4, idx, spec.header)
    for r, row in enumerate(rows, start=5):
        for idx, spec in enumerate(columns, start=1):
            value = _excel_text(row[spec.key]) if spec.number_format == "@" else _normalize(row[spec.key])
            _put(sheet, r, idx, value)

    _style(
        sheet, f"A4:{end_letter}4", fill=ESPRESSO, font=_font(WHITE, 9, bold=True), horizontal="center", vertical="center", wrap=True,
    )
    sheet.row_dimensions[4].height = 34
    if rows:
        _style(sheet, f"A5:{end_letter}{end_row}", font=_font(TEXT, 9), vertical="center")
        _row_height(sheet, 5, end_row, row_height)

    wrap = set(wrap_keys)
    jsons = set(json_keys)
    center_keys = {
        "warehouse_type", "warehouse_size", "min_clusters", "max_clusters", "warehouse_channel", "photon_status",
        "requires_hitl", "scenario_confidence_level", "evidence_confidence_level", "status", "recommended_flag",
        "confidence_level", "compatibility_status", "remeasure_required", "economics_available", "is_additive",
        "realization_status", "validation_status", "rollback_required", "knob", "package_base_impact_type",
        "base_impact_type", "recommended_impact_type",
    }

    for idx, spec in enumerate(columns, start=1):
        letter = _col_letter(idx)
        sheet.column_dimensions[letter].width = spec.width
        data_ref = f"{letter}5:{letter}{end_row}"
        if not rows:
            continue
        if spec.number_format:
            _style(sheet, data_ref, number_format=spec.number_format)
        elif spec.key.endswith("_json") or spec.key == "package_file":
            _style(sheet, data_ref, number_format="@")
        if spec.key in wrap:
            _style(sheet, data_ref, wrap=True, vertical="top")
        if spec.key in center_keys:
            _style(sheet, data_ref, horizontal="center")
        if spec.key in jsons:
            _style(sheet, data_ref, font=_font(MUTED, 8), wrap=False)
        if spec.number_format:
            _style(sheet, data_ref, horizontal="right")

    if rows:
        _add_table(sheet, f"A4:{end_letter}{end_row}", table_name, "TableStyleMedium2")

    sheet.freeze_panes = f"{_col_letter(freeze_columns + 1)}5"
    _apply_status_formatting(sheet, {c.key: i + 1 for i, c in enumerate(columns)}, 5, end_row)


def _add_fill_rule(sheet: Any, ref: str, formula: str, color: str) -> None:
    sheet.conditional_formatting.add(ref, FormulaRule(formula=[formula], fill=_fill(color)))


def _apply_status_formatting(sheet: Any, key_to_col: Mapping[str, int], start_row: int, end_row: int) -> None:
    if end_row < start_row:
        return
    for key in (
        "status", "overall_status", "evidence_quality_status", "scenario_confidence_level", "evidence_confidence_level",
        "confidence_level", "realization_status", "validation_status", "compatibility_status",
        "package_base_impact_type", "base_impact_type", "recommended_impact_type",
    ):
        col = key_to_col.get(key)
        if not col:
            continue
        letter = _col_letter(col)
        ref = f"{letter}{start_row}:{letter}{end_row}"
        for word, fill in [
            ("ACTIONABLE", GREEN), ("RECOMMENDED", GREEN), ("HIGH", GREEN), ("SUFFICIENT", GREEN), ("PASS", GREEN), ("REALIZED", GREEN),
            ("TEST_REQUIRED", AMBER), ("REVIEW", AMBER), ("MEDIUM", AMBER), ("PARTIAL", AMBER), ("REQUIRES_VALIDATION", AMBER), ("VALIDATING", AMBER),
            ("BLOCKED", RED), ("LOW", RED), ("INSUFFICIENT_EVIDENCE", RED), ("FAILED_VALIDATION", RED), ("ROLLED_BACK", RED),
            ("SAVINGS", GREEN), ("COST_INCREASE", RED), ("MIXED", AMBER),
            ("NO_CHANGE", GRAY), ("NOT_APPLICABLE", GRAY), ("ALTERNATIVE", BLUE), ("UNAVAILABLE", GRAY), ("NOT_APPLIED", GRAY), ("UNKNOWN", GRAY),
        ]:
            _add_fill_rule(sheet, ref, f'{letter}{start_row}="{word}"', fill)

    # Boolean flags that matter operationally.
    for key, fill_true in (("requires_hitl", AMBER), ("recommended_flag", GREEN), ("remeasure_required", AMBER), ("rollback_required", RED)):
        col = key_to_col.get(key)
        if col:
            letter = _col_letter(col)
            _add_fill_rule(sheet, f"{letter}{start_row}:{letter}{end_row}", f"{letter}{start_row}=TRUE", fill_true)


def _build_executive(wb: Any, projection: wc.WorkbookProjection) -> None:
    sheet = wb.create_sheet(wc.SHEET_EXECUTIVE)
    _sheet_title(
        sheet,
        "SQL Warehouse Optimization — Executive Summary",
        "Decision-ready portfolio view from Reporter JSON Contract v1.1.0. Financial cells are numeric; impact direction is explicit and sortable.",
        20,
    )
    # Section labels make the side-by-side layout obvious.
    sheet.merge_cells("A3:E3")
    sheet["A3"] = "Portfolio metrics"
    _style(sheet, "A3", fill=GOLD, font=_font(WHITE, 9, bold=True), horizontal="center")
    sheet.merge_cells("G3:T3")
    sheet["G3"] = "Top warehouse impacts (up to 15)"
    _style(sheet, "G3", fill=ESPRESSO_2, font=_font(WHITE, 9, bold=True), horizontal="center")
    sheet.row_dimensions[3].height = 20

    metrics = projection.executive_metrics
    metric_cols = [("section", "Section"), ("metric", "Metric"), ("value", "Value"), ("currency", "Currency"), ("notes", "Notes")]
    metric_end = 4 + len(metrics)
    for j, (_, header) in enumerate(metric_cols, start=1):
        _put(sheet, 4, j, header)
    for r, row in enumerate(metrics, start=5):
        for j, (key, _) in enumerate(metric_cols, start=1):
            _put(sheet, r, j, row[key])
    _style(sheet, "A4:E4", fill=ESPRESSO, font=_font(WHITE, 9, bold=True), wrap=True)
    sheet.row_dimensions[4].height = 30
    for i, w in enumerate([16, 30, 34, 10, 45], 1):
        sheet.column_dimensions[_col_letter(i)].width = w
    if metrics:
        _style(sheet, f"E5:E{metric_end}", wrap=True)
        _row_height(sheet, 5, metric_end, 23)
    for i, row in enumerate(metrics, start=5):
        value = row.get("value")
        if row.get("currency") not in (wc.DISPLAY_NA, "MIXED") and isinstance(value, (int, float)) and not isinstance(value, bool):
            sheet[f"C{i}"].number_format = MONEY_FMT
        elif isinstance(value, int) and not isinstance(value, bool):
            sheet[f"C{i}"].number_format = INT_FMT
    if metrics:
        _add_table(sheet, f"A4:E{metric_end}", "ExecutiveMetrics", "TableStyleMedium2")

    opp_cols = [
        ("warehouse_name", "Warehouse"), ("owner", "Owner"), ("currency", "Currency"),
        ("forward_total_cost_365d", "Forward baseline"), ("recommended_impact_type", "Impact type"),
        ("recommended_annual_impact", "Recommended annual impact"), ("scenario_low_impact", "Scenario low"),
        ("scenario_high_impact", "Scenario high"), ("impact_pct", "Impact %"),
        ("recommended_scenario_name", "Recommended plan"), ("scenario_confidence_level", "Confidence"),
        ("next_action", "Next action"), ("savings_amount", "Savings"), ("cost_increase_amount", "Cost increase"),
    ]
    opp_rows = projection.executive_opportunities
    opp_end = 4 + len(opp_rows)
    for j, (_, header) in enumerate(opp_cols, start=7):
        _put(sheet, 4, j, header)
    for r, row in enumerate(opp_rows, start=5):
        for j, (key, _) in enumerate(opp_cols, start=7):
            _put(sheet, r, j, row[key])
    _style(sheet, "G4:T4", fill=ESPRESSO, font=_font(WHITE, 9, bold=True), wrap=True)
    sheet.row_dimensions[4].height = 30
    for j, width in enumerate([22, 32, 9, 16, 15, 18, 15, 15, 11, 34, 12, 30, 14, 14], start=7):
        sheet.column_dimensions[_col_letter(j)].width = width
    if opp_rows:
        for c in ("J", "L", "M", "N", "S", "T"):
            _style(sheet, f"{c}5:{c}{opp_end}", number_format=MONEY_FMT)
        _style(sheet, f"O5:O{opp_end}", number_format=PCT_FMT)
        _row_height(sheet, 5, opp_end, 30)
        for c in ("G", "P", "R"):
            _style(sheet, f"{c}5:{c}{opp_end}", wrap=True)
        _style(sheet, f"H5:H{opp_end}", wrap=False)
        _add_table(sheet, f"G4:T{opp_end}", "MaterialImpacts", "TableStyleMedium4")
        for col, word, fill in (("K", "SAVINGS", GREEN), ("K", "COST_INCREASE", RED), ("K", "NO_CHANGE", GRAY), ("Q", "HIGH", GREEN), ("Q", "MEDIUM", AMBER), ("Q", "LOW", RED)):
            _add_fill_rule(sheet, f"{col}5:{col}{opp_end}", f'{col}5="{word}"', fill)
        chart = BarChart()
        chart.type = "col"
        chart.title = "Recommended annual economic impact"
        chart.add_data(Reference(sheet, min_col=19, max_col=20, min_row=4, max_row=opp_end), titles_from_data=True)  # S:T
        chart.set_categories(Reference(sheet, min_col=7, min_row=5, max_row=opp_end))  # G
        chart.y_axis.number_format = "#,##0"
        chart.legend.position = "r"
        chart.width, chart.height = 32, 9.5
        sheet.add_chart(chart, "G22")
    sheet.freeze_panes = "A5"


def _build_portfolio(wb: Any, projection: wc.WorkbookProjection) -> None:
    sheet = wb.create_sheet(wc.SHEET_PORTFOLIO)
    _sheet_title(sheet, "Warehouse Portfolio", "One row per recommendation package. All agreed contract columns are retained; verbose audit fields are placed at the far right.", len(wc.PORTFOLIO_COLUMNS))
    _write_contract_table(
        sheet, columns=wc.PORTFOLIO_COLUMNS, groups=wc.PORTFOLIO_GROUPS, rows=projection.warehouse_portfolio,
        table_name="WarehousePortfolio", freeze_columns=4, row_height=24,
        wrap_keys=("warehouse_name", "recommended_scenario_name", "next_action"),
        json_keys=("query_source_profile_json", "scenario_ids_json", "candidate_ids_json", "package_file", "evidence_summary_json"),
    )


def _build_scenarios(wb: Any, projection: wc.WorkbookProjection) -> None:
    sheet = wb.create_sheet(wc.SHEET_SCENARIOS)
    _sheet_title(sheet, "Scenario Plans", "One row per scenario. Economics are numeric and sortable; signed annual impact is positive for savings and negative for cost increase.", len(wc.SCENARIO_COLUMNS))
    _write_contract_table(
        sheet, columns=wc.SCENARIO_COLUMNS, groups=wc.SCENARIO_GROUPS, rows=projection.scenario_plans,
        table_name="ScenarioPlans", freeze_columns=5, row_height=24,
        wrap_keys=("warehouse_name", "scenario_name"),
        json_keys=("compatibility_checks_json", "compatibility_blockers_json", "execution_plan_json", "candidate_ids_json", "dependencies_json", "invalidations_json"),
    )


def _build_registry(wb: Any, projection: wc.WorkbookProjection) -> None:
    sheet = wb.create_sheet(wc.SHEET_REGISTRY)
    _sheet_title(sheet, "Recommendation Registry", "One row per candidate. Decision, economics, realization and audit evidence align to Workbook Contract v1.2.0.", len(wc.REGISTRY_COLUMNS))
    _write_contract_table(
        sheet, columns=wc.REGISTRY_COLUMNS, groups=wc.REGISTRY_GROUPS, rows=projection.recommendation_registry,
        table_name="RecommendationRegistry", freeze_columns=6, row_height=32,
        wrap_keys=("warehouse_name", "objectives", "current_value", "recommended_value", "rationale", "decision_rule"),
        json_keys=("confidence_components_json", "evidence_refs_json", "evidence_json", "thresholds_json", "dependencies_json", "invalidates_json", "blockers_json", "scenario_membership_json", "realization_records_json"),
    )


def write_workbook(path: str | Path, packages: Sequence[rc.RecommendationPackage]) -> Path:
    """Generate the locked four-sheet workbook projection and save it to *path*."""
    for package in packages:
        package.validate()
    projection = wc.project_workbook(packages)
    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet: the contract has exactly four visible sheets
    _build_executive(wb, projection)
    _build_portfolio(wb, projection)
    _build_scenarios(wb, projection)
    _build_registry(wb, projection)
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(target))
    return target
