"""The realization workbook renderer (Realization Workbook Contract ``report.REALIZATION_WORKBOOK_CONTRACT_VERSION``).

Lays out ``report.project_realization_workbook`` with ``openpyxl``: three visible sheets (Warehouses, Lineages, Timeline),
each a group banner row, a header row and one row per projected record. The sheets, column order, headers, groups and the
kind of each column all come from ``realizer.contracts.report``; this module keeps no second list and decides, infers and
computes nothing. The rules are the Publisher workbook's, implemented here so that the products stay independent: no
formulas, macros or hidden sheets; identifiers written as text; money as signed numbers next to a currency column;
unknown values as empty cells, never zero; audit JSON at the far right; the header row frozen.

The file's document properties (created, modified) are the latest time the inputs were observed, so identical inputs
give identical content.
"""
from __future__ import annotations

import datetime as dt
import zipfile
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.writer.excel import ExcelWriter

from realizer.common.utils import parse_utc
from realizer.contracts import report

GROUP_ROW = 1
HEADER_ROW = 2
FIRST_DATA_ROW = HEADER_ROW + 1

MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00;0.00"  # signed: a cost increase stays negative
RATIO_FORMAT = "0.00"
NUMBER_FORMAT = "#,##0.00"
COUNT_FORMAT = "#,##0"
TEXT_FORMAT = "@"
FORMATS = {
    report.KIND_ID: TEXT_FORMAT, report.KIND_TIME: TEXT_FORMAT, report.KIND_JSON: TEXT_FORMAT,
    report.KIND_MONEY: MONEY_FORMAT, report.KIND_RATIO: RATIO_FORMAT, report.KIND_NUMBER: NUMBER_FORMAT, report.KIND_COUNT: COUNT_FORMAT,
}
TEXT_KINDS = {report.KIND_ID, report.KIND_TIME, report.KIND_JSON, report.KIND_TEXT}

# Columns frozen at the left of each sheet, with the header rows.
FROZEN_COLUMNS = {report.SHEET_WAREHOUSES: 3, report.SHEET_LINEAGES: 2, report.SHEET_TIMELINE: 1}
TABLE_NAMES = {report.SHEET_WAREHOUSES: "Warehouses", report.SHEET_LINEAGES: "Lineages", report.SHEET_TIMELINE: "Timeline"}

# Excel rejects longer cells; an audit JSON cell is cut with a visible marker (the full content is in the RPKG and ledger).
EXCEL_MAX_CELL_CHARS = 32767
TRUNCATION_MARKER = " ...[truncated; full content in the realization package and ledger]"
# The time a workbook of no observed input is stamped with.
NO_OBSERVATION_TIME = dt.datetime(2000, 1, 1)

HEADER_FILL = PatternFill(fill_type="solid", start_color="3B2F2F", end_color="3B2F2F")
GROUP_FILLS = (PatternFill(fill_type="solid", start_color="B58A3A", end_color="B58A3A"),
               PatternFill(fill_type="solid", start_color="5A463B", end_color="5A463B"))
WHITE_BOLD = Font(color="FFFFFF", size=9, bold=True)
BODY_FONT = Font(color="2B2522", size=9)
AUDIT_FONT = Font(color="6B625D", size=8)


def _cell_value(spec: report.ColumnSpec, value: Any) -> Any:
    """The value written for *spec*: None stays an empty cell; text kinds are written as text (an id never becomes a
    number); a number stays a number."""
    if value is None:
        return None
    if spec.kind in TEXT_KINDS:
        text = str(value)
        if len(text) > EXCEL_MAX_CELL_CHARS:
            text = text[: EXCEL_MAX_CELL_CHARS - len(TRUNCATION_MARKER)] + TRUNCATION_MARKER
        return text
    return value


def _put(sheet: Any, row: int, column: int, value: Any) -> Any:
    """Write one cell; a text is never read as a formula."""
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _write_sheet(sheet: Any, name: str, rows: Sequence[Mapping[str, Any]]) -> None:
    columns, groups = report.SHEET_COLUMNS[name]
    position = {c.key: i for i, c in enumerate(columns, start=1)}
    for i, group in enumerate(groups):
        start, end = position[group.start_key], position[group.end_key]
        cell = _put(sheet, GROUP_ROW, start, group.title)
        cell.fill, cell.font = GROUP_FILLS[i % 2], WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if end > start:
            sheet.merge_cells(start_row=GROUP_ROW, start_column=start, end_row=GROUP_ROW, end_column=end)
    for i, spec in enumerate(columns, start=1):
        cell = _put(sheet, HEADER_ROW, i, spec.header)
        cell.fill, cell.font = HEADER_FILL, WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(i)].width = spec.width
    sheet.row_dimensions[HEADER_ROW].height = 30
    for r, row in enumerate(rows, start=FIRST_DATA_ROW):
        for i, spec in enumerate(columns, start=1):
            cell = _put(sheet, r, i, _cell_value(spec, row[spec.key]))
            cell.font = AUDIT_FONT if spec.kind == report.KIND_JSON else BODY_FONT
            if spec.kind in FORMATS:
                cell.number_format = FORMATS[spec.kind]
    if rows:
        ref = f"A{HEADER_ROW}:{get_column_letter(len(columns))}{HEADER_ROW + len(rows)}"
        table = Table(displayName=TABLE_NAMES[name], ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    sheet.freeze_panes = f"{get_column_letter(FROZEN_COLUMNS[name] + 1)}{FIRST_DATA_ROW}"


def write_realization_workbook(path: Path, warehouse_runs: Sequence[report.WarehouseRun]) -> Path:
    """Write the realization workbook of *warehouse_runs* to *path* and return the path written.

    The rows are ``report.project_realization_workbook``'s (which validates them first: nothing is written when it
    fails). Creates parent directories and overwrites an existing file.
    """
    projection = report.project_realization_workbook(warehouse_runs)
    book = Workbook()
    book.remove(book.active)  # exactly the contract's sheets
    for name, rows in zip(report.SHEET_NAMES, (projection.warehouses, projection.lineages, projection.timeline)):
        _write_sheet(book.create_sheet(name), name, rows)
    stamp = (parse_utc(projection.as_of_utc).replace(tzinfo=None) if projection.as_of_utc else NO_OBSERVATION_TIME)
    book.properties.created = stamp
    book.properties.modified = stamp
    book.properties.creator = "Realizer"
    book.properties.lastModifiedBy = "Realizer"
    book.properties.title = f"Realization workbook (contract {report.REALIZATION_WORKBOOK_CONTRACT_VERSION})"
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    # openpyxl's ``Workbook.save`` stamps ``modified`` with the wall clock; its writer, given the archive, keeps the stamp above.
    ExcelWriter(book, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, allowZip64=True)).save()
    return target
