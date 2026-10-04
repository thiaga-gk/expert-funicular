import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import publisher as p


class PublisherV27Tests(unittest.TestCase):
    def setUp(self):
        self.t = p.Thresholds()
        self.evidence = p.demo_evidence()
        self.results = p.evaluate_warehouses(self.evidence, self.t)
        self.w = self.results[0]
        self.by_knob = {c.knob: c for c in self.w.candidates}
        self.packages = p.build_recommendation_packages(self.evidence, self.t)

    def test_six_deterministic_candidates_and_no_photon_recommendation(self):
        self.assertEqual(set(self.by_knob), set(p.ACTIVE_KNOBS))
        self.assertEqual(len(self.w.candidates), 6)
        self.assertNotIn("enable_photon", self.by_knob)

    def test_expected_demo_decisions(self):
        self.assertEqual(self.by_knob["warehouse_size"].status, "TEST_REQUIRED")
        self.assertEqual(self.by_knob["warehouse_size"].recommended_config, {"warehouse_size": "MEDIUM"})
        self.assertEqual(self.by_knob["max_clusters"].recommended_config, {"max_clusters": 3})
        self.assertEqual(self.by_knob["min_clusters"].recommended_config, {"min_clusters": 1})
        self.assertEqual(self.by_knob["auto_stop_minutes"].recommended_config, {"auto_stop_minutes": 30})
        self.assertEqual(self.by_knob["warehouse_type"].recommended_config, {"warehouse_type": "SERVERLESS"})

    def test_recommendation_ids_are_deterministic_and_include_current_config(self):
        r2 = p.evaluate_warehouses(p.demo_evidence(), self.t)
        self.assertEqual([c.recommendation_id for c in self.w.candidates], [c.recommendation_id for c in r2[0].candidates])
        old = self.by_knob["warehouse_size"].recommendation_id
        changed = p.demo_evidence()
        changed[0].current_epoch.config = p.WarehouseConfig("PRO", "CURRENT", "X_LARGE", 2, 4, 120)
        new = next(c for c in p.evaluate_warehouses(changed, self.t)[0].candidates if c.knob == "warehouse_size").recommendation_id
        self.assertNotEqual(old, new)

    def test_analysis_time_makes_decision_semantics_repeatable(self):
        first = p.evaluate_warehouses(p.demo_evidence(), self.t)[0]
        second = p.evaluate_warehouses(p.demo_evidence(), self.t)[0]
        self.assertEqual(p.canonical_json(first.candidates), p.canonical_json(second.candidates))
        self.assertEqual(p.canonical_json(first.scenarios), p.canonical_json(second.scenarios))

    def test_current_epoch_drives_findings_and_history_calibrates_estimator(self):
        current = self.evidence[0].current_epoch_id
        for c in self.w.candidates:
            for f in c.findings:
                if f.metric != "serverless_usage_pct_365":
                    self.assertEqual(f.epoch_id, current)
        size = self.by_knob["warehouse_size"]
        calib = size.metadata.get("historical_epoch_calibration")
        self.assertIsNotNone(calib)
        self.assertNotEqual(calib["epoch_id"], current)

    def test_candidates_are_independent_and_portfolio_is_not_sum(self):
        standalone = sum((c.estimated_savings_365 or 0) for c in self.w.candidates)
        chosen = next(s for s in self.w.scenarios if s.recommended)
        portfolio = chosen.portfolio_estimate.estimated_savings_range_365
        self.assertIsNotNone(portfolio)
        self.assertNotAlmostEqual(standalone, portfolio.base, places=2)
        self.assertTrue(chosen.portfolio_estimate.non_additive)
        self.assertTrue(all(c.metadata["candidate_savings_must_not_be_summed"] for c in self.w.candidates))

    def test_stateful_bundle_reduces_overlap_between_min_and_autostop(self):
        chosen = next(s for s in self.w.scenarios if s.recommended)
        portfolio = chosen.portfolio_estimate.estimated_savings_range_365
        independent_high = sum(
            self.by_knob[k].estimated_savings_range_365.high
            for k in ("warehouse_size", "min_clusters", "auto_stop_minutes")
            if self.by_knob[k].estimated_savings_range_365
        )
        self.assertLess(portfolio.high, independent_high)
        self.assertLessEqual(portfolio.high, self.w.total_365d_baseline_cost)
        self.assertIn("cluster-minutes", chosen.portfolio_estimate.method)

    def test_min_cluster_reduction_is_a_canary_candidate_not_actionable(self):
        # H4: time at the floor proves the autoscaler sat at min, not that a floor cluster was spare.
        c = self.by_knob["min_clusters"]
        self.assertEqual(c.status, "TEST_REQUIRED")
        self.assertIn("SPARE_FLOOR_CAPACITY_UNPROVEN", c.blockers)
        self.assertEqual(c.recommended_config, {"min_clusters": 1})
        plan = next(s for s in self.w.scenarios if s.recommended)
        step = next(x for x in plan.steps if x.knob == "min_clusters")
        self.assertIn("CANARY", step.validation_gate)

    def test_min_cluster_time_at_min_is_an_upper_bound_not_the_base(self):
        c = self.by_knob["min_clusters"]
        r = c.estimated_savings_range_365
        # Low: the floor capacity may have been needed, so the change may save nothing.
        self.assertAlmostEqual(r.low, 0.0, places=2)
        # High is the old time-at-min figure, now labelled as an upper bound; base is the policy midpoint of it.
        self.assertGreater(r.high, 0)
        self.assertAlmostEqual(r.base, r.high * self.t.min_cluster_realization_base, places=2)
        self.assertLess(r.base, 0.75 * r.high)  # the pre-fix base was 75% of this figure
        self.assertTrue(c.metadata["savings_high_is_upper_bound"])
        self.assertIn("upper bound", c.metadata["method"])

    def test_bundle_prices_min_cluster_with_the_same_realization_as_the_candidate(self):
        c = self.by_knob["min_clusters"]
        step = p.PlanStep(1, c.recommendation_id, "min_clusters", c.current_config, c.recommended_config,
                          c.status, "A_B_CANARY", 7, [], [])
        port = p.StatefulBundleEstimator(self.t).estimate(self.evidence[0], [step], {c.recommendation_id: c})
        got, want = port.estimated_savings_range_365, c.estimated_savings_range_365
        for g, w in ((got.low, want.low), (got.base, want.base), (got.high, want.high)):
            self.assertAlmostEqual(g, w, delta=1.0)

    # --- audit H1: steps marked PENDING_REEVALUATION must not feed authoritative scenario economics ----------
    def plan_a(self):
        return next(s for s in self.w.scenarios if s.name.startswith("Plan A"))

    def test_plan_a_authoritative_savings_stop_at_the_first_invalidating_step(self):
        plan = self.plan_a()
        states = {x.knob: x.execution_state for x in plan.steps}
        self.assertEqual(states["warehouse_size"], "READY_FOR_CANARY")
        for knob in ("max_clusters", "min_clusters", "auto_stop_minutes"):
            self.assertEqual(states[knob], "PENDING_REEVALUATION")
        size = next(x for x in plan.steps if x.knob == "warehouse_size")
        auth = plan.portfolio_estimate.estimated_savings_range_365
        # The authoritative scenario figure is exactly the size step; pending steps add nothing to it.
        self.assertAlmostEqual(auth.base, size.running_estimated_savings_range_365.base, places=2)
        self.assertAlmostEqual(auth.high, size.running_estimated_savings_range_365.high, places=2)

    def test_pending_steps_are_priced_as_a_labeled_contingent_upside(self):
        plan = self.plan_a()
        basis = {x.knob: x.pricing_basis for x in plan.steps}
        self.assertEqual(basis["warehouse_size"], "AUTHORITATIVE")
        for knob in ("max_clusters", "min_clusters", "auto_stop_minutes"):
            self.assertEqual(basis[knob], "CONTINGENT_PENDING_REEVALUATION")
        contingent = plan.portfolio_estimate.contingent
        self.assertIsNotNone(contingent)
        last = plan.steps[-1]
        self.assertAlmostEqual(contingent.estimated_savings_range_365.base, last.running_estimated_savings_range_365.base, places=2)
        # The big-picture upside is larger than the authoritative base, and clearly separate from it.
        self.assertGreater(contingent.estimated_savings_range_365.base, plan.portfolio_estimate.estimated_savings_range_365.base)
        self.assertEqual(contingent.estimate_status, "PROVISIONAL_MODELED_RANGE")

    def step_for(self, knob, state):
        c = self.by_knob[knob]
        return p.PlanStep(1, c.recommendation_id, knob, c.current_config, c.recommended_config, c.status,
                          "GATE", 0, c.invalidates, c.dependencies, state), c

    def test_authoritative_chain_skips_pending_effects_even_when_a_later_step_is_ready(self):
        pending, pc = self.step_for("min_clusters", "PENDING_REEVALUATION")
        ready, rcand = self.step_for("auto_stop_minutes", "READY_FOR_HITL")
        by_id = {pc.recommendation_id: pc, rcand.recommendation_id: rcand}
        est = p.StatefulBundleEstimator(self.t)
        both = est.estimate(self.evidence[0], [pending, ready], by_id)
        only_ready = est.estimate(self.evidence[0], [self.step_for("auto_stop_minutes", "READY_FOR_HITL")[0]], by_id)
        # Authoritative: auto-stop is priced on the *unchanged* floor because the min-cluster step is pending.
        self.assertAlmostEqual(both.estimated_savings_range_365.base, only_ready.estimated_savings_range_365.base, places=2)
        # Contingent: the full chain, where min-cluster consumed minutes first, so auto-stop has less left to remove.
        self.assertIsNotNone(both.contingent)
        self.assertNotAlmostEqual(both.contingent.estimated_savings_range_365.base, both.estimated_savings_range_365.base, places=2)

    def test_no_contingent_economics_when_nothing_is_pending(self):
        ready, c = self.step_for("auto_stop_minutes", "READY_FOR_HITL")
        port = p.StatefulBundleEstimator(self.t).estimate(self.evidence[0], [ready], {c.recommendation_id: c})
        self.assertIsNone(port.contingent)
        self.assertEqual(ready.pricing_basis, "AUTHORITATIVE")
        plan_b = next(s for s in self.w.scenarios if s.name.startswith("Plan B"))
        self.assertIsNone(plan_b.portfolio_estimate.contingent)

    def test_unpriceable_pending_step_does_not_poison_authoritative_economics(self):
        ready, rcand = self.step_for("auto_stop_minutes", "READY_FOR_HITL")
        # A pending step whose effect only a canary can measure withholds only the contingent chain.
        photon = p.PlanStep(2, "REC-UNMEASURABLE", "max_clusters", {"max_clusters": 4}, {"max_clusters": 3},
                            "TEST_REQUIRED", "GATE", 0, [], [], "PENDING_REEVALUATION")
        photon_c = copy.deepcopy(self.by_knob["max_clusters"])
        photon_c.recommendation_id = "REC-UNMEASURABLE"
        photon_c.metadata = {**photon_c.metadata, "bundle_effect": {"type": "requires_canary"}}
        by_id = {rcand.recommendation_id: rcand, "REC-UNMEASURABLE": photon_c}
        port = p.StatefulBundleEstimator(self.t).estimate(self.evidence[0], [ready, photon], by_id)
        only_ready = p.StatefulBundleEstimator(self.t).estimate(self.evidence[0], [self.step_for("auto_stop_minutes", "READY_FOR_HITL")[0]], by_id)
        self.assertEqual(port.estimate_status, only_ready.estimate_status)
        self.assertAlmostEqual(port.estimated_savings_range_365.base, only_ready.estimated_savings_range_365.base, places=2)
        self.assertIsNone(port.contingent)  # contingent cannot be priced, so it is withheld rather than guessed
        self.assertIsNone(photon.running_estimated_savings_range_365)

    def test_reviewer_gate_2_rejects_pending_steps_priced_as_authoritative(self):
        w = copy.deepcopy(self.w)
        plan = next(s for s in w.scenarios if s.recommended)
        next(x for x in plan.steps if x.execution_state == "PENDING_REEVALUATION").pricing_basis = "AUTHORITATIVE"
        with self.assertRaisesRegex(ValueError, "pricing basis"):
            p.Reviewer(self.t).review_scenarios(self.evidence[0], w.candidates, w.scenarios, plan.scenario_id)

    def test_reviewer_gate_2_rejects_contingent_upside_leaking_into_authoritative_savings(self):
        w = copy.deepcopy(self.w)
        plan = next(s for s in w.scenarios if s.recommended)
        plan.portfolio_estimate.estimated_savings_range_365 = plan.portfolio_estimate.contingent.estimated_savings_range_365
        with self.assertRaisesRegex(ValueError, "Authoritative scenario savings"):
            p.Reviewer(self.t).review_scenarios(self.evidence[0], w.candidates, w.scenarios, plan.scenario_id)

    def test_reviewer_gate_2_accepts_the_unmodified_plan(self):
        plan = next(s for s in self.w.scenarios if s.recommended)
        p.Reviewer(self.t).review_scenarios(self.evidence[0], self.w.candidates, self.w.scenarios, plan.scenario_id)

    def test_scenario_contract_and_order(self):
        self.assertEqual([s.name for s in self.w.scenarios], [
            "Plan A — Optimize current warehouse type", "Plan B — Serverless migration experiment"
        ])
        plan_a = next(s for s in self.w.scenarios if s.name.startswith("Plan A"))
        plan_b = next(s for s in self.w.scenarios if s.name.startswith("Plan B"))
        self.assertTrue(plan_a.recommended)
        self.assertEqual([x.knob for x in plan_a.steps], ["warehouse_size", "max_clusters", "min_clusters", "auto_stop_minutes"])
        self.assertEqual(plan_b.portfolio_estimate.estimate_status, "PRE_CANARY_MODELED_RANGE")
        self.assertIsNotNone(plan_b.portfolio_estimate.dbu_365d_savings_range)
        self.assertIsNotNone(plan_b.portfolio_estimate.aws_365d_savings_range)
        self.assertEqual([x.knob for x in self.w.recommended_execution_plan], [x.knob for x in plan_a.steps])

    def test_portfolio_range_ordering_and_plan_running_ranges(self):
        plan = next(s for s in self.w.scenarios if s.recommended)
        r = plan.portfolio_estimate.estimated_savings_range_365
        self.assertLessEqual(r.low, r.base)
        self.assertLessEqual(r.base, r.high)
        self.assertTrue(all(s.running_estimated_savings_range_365 is not None for s in plan.steps))

    def test_reporter_json_workbook_and_hitl_findings(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            packages = self.packages
            files = p.package_locked_contract(Path(td), packages, self.evidence)
            payload = json.loads(Path(files[f"package:{self.evidence[0].workspace_id}:{self.evidence[0].warehouse_id}"]).read_text())
            for field in ("candidates", "scenarios", "recommended_plan", "package_economics", "evidence", "workload_profile", "realization"):
                self.assertIn(field, payload)
            self.assertEqual(len(payload["candidates"]), 6)
            self.assertGreater(len(payload["evidence"]), 0)
            wb = load_workbook(files["workbook"])
            self.assertEqual(wb.sheetnames, ["Executive Summary", "Warehouse Portfolio", "Scenario Plans", "Recommendation Registry"])
    def test_read_only_api_scope_and_supported_type_gate_are_explicit(self):
        self.assertEqual(p.SUPPORTED_WAREHOUSE_TYPES, {"CLASSIC", "PRO", "SERVERLESS"})
        source = Path(p.__file__).read_text()
        self.assertIn("/api/2.0/sql/config/warehouses", source)
        self.assertIn("/api/2.1/clusters/get", source)
        self.assertNotIn("/api/2.0/sql/warehouses/{id}", source)

    def _fleet_with_one_unsupported(self):
        fleet = p.synthetic_demo_evidence(3)
        fleet[1].current_epoch.config = p.WarehouseConfig("LAKEHOUSE_REAL_TIME", "CURRENT", "MEDIUM", 1, 1, 10)
        return fleet

    def test_unsupported_type_is_excluded_per_warehouse_not_whole_run(self):
        # G1: one unsupported warehouse must not abort the run for every other warehouse.
        fleet = self._fleet_with_one_unsupported()
        packages = p.build_recommendation_packages(fleet, self.t)
        keys = {(pkg.warehouse.workspace_id, pkg.warehouse.warehouse_id) for pkg in packages}
        self.assertEqual(len(packages), 2)
        self.assertNotIn((fleet[1].workspace_id, fleet[1].warehouse_id), keys)
        supported_only = [e for i, e in enumerate(p.synthetic_demo_evidence(3)) if i != 1]
        expected = p.build_recommendation_packages(supported_only, self.t)
        self.assertEqual(
            sorted(pkg.recommendation_package_id for pkg in packages),
            sorted(pkg.recommendation_package_id for pkg in expected),
        )

    def test_unsupported_type_never_reaches_the_analyzer(self):
        seen = []
        original = p.Analyzer.analyze

        def spy(analyzer, e):
            seen.append(e.current_epoch.config.warehouse_type)
            return original(analyzer, e)

        with patch.object(p.Analyzer, "analyze", spy):
            p.evaluate_warehouses(self._fleet_with_one_unsupported(), self.t)
        self.assertTrue(seen)
        self.assertTrue(set(seen) <= p.SUPPORTED_WAREHOUSE_TYPES)

    def test_unsupported_warehouse_does_not_feed_workspace_calibration(self):
        fleet = p.synthetic_demo_evidence(2)
        unsupported = fleet[1]
        prior = copy.deepcopy(unsupported.current_epoch)
        prior.is_current = False
        prior.epoch_id = prior.epoch_id + "-prior"
        prior.config = p.WarehouseConfig("SERVERLESS", "CURRENT", "MEDIUM", 1, 1, 10)
        unsupported.epochs.insert(0, prior)
        unsupported.current_epoch.config = p.WarehouseConfig("FUTURE_TYPE", "CURRENT", "MEDIUM", 1, 1, 10)
        self.assertFalse(any(ep.config.warehouse_type == "SERVERLESS" for ep in fleet[0].epochs))
        p.evaluate_warehouses(fleet, self.t)
        self.assertFalse(fleet[0].billing_features["workspace_serverless_observed_from_system_compute_warehouses"])

    def test_unsupported_exclusion_recorded_in_manifest(self):
        fleet = self._fleet_with_one_unsupported()
        packages = p.build_recommendation_packages(fleet, self.t)
        with tempfile.TemporaryDirectory() as td:
            files = p.package_locked_contract(Path(td), packages, fleet, generate_workbook=False)
            manifest = json.loads(Path(files["manifest"]).read_text())
        self.assertEqual(manifest["package_count"], 2)
        self.assertEqual(manifest["excluded_warehouse_count"], 1)
        self.assertEqual(manifest["excluded_warehouses"], [{
            "workspace_id": fleet[1].workspace_id,
            "warehouse_id": fleet[1].warehouse_id,
            "warehouse_name": fleet[1].warehouse_name,
            "warehouse_type": "LAKEHOUSE_REAL_TIME",
            "reason_code": "UNSUPPORTED_WAREHOUSE_TYPE",
        }])

    def test_all_unsupported_fleet_emits_no_packages_and_records_every_exclusion(self):
        fleet = p.synthetic_demo_evidence(2)
        for e in fleet:
            e.current_epoch.config = p.WarehouseConfig("FUTURE_TYPE", "CURRENT", "MEDIUM", 1, 1, 10)
        packages = p.build_recommendation_packages(fleet, self.t)
        self.assertEqual(packages, [])
        with tempfile.TemporaryDirectory() as td:
            files = p.package_locked_contract(Path(td), packages, fleet, generate_workbook=False)
            manifest = json.loads(Path(files["manifest"]).read_text())
        self.assertEqual(manifest["package_count"], 0)
        self.assertEqual([x["warehouse_type"] for x in manifest["excluded_warehouses"]], ["FUTURE_TYPE", "FUTURE_TYPE"])

    def test_supported_manifest_has_empty_exclusions(self):
        with tempfile.TemporaryDirectory() as td:
            files = p.package_locked_contract(Path(td), self.packages, self.evidence, generate_workbook=False)
            manifest = json.loads(Path(files["manifest"]).read_text())
        self.assertEqual(manifest["excluded_warehouse_count"], 0)
        self.assertEqual(manifest["excluded_warehouses"], [])

    def test_query_source_tag_cache_evidence_exposed(self):
        metrics = {f.metric for c in self.w.candidates for f in c.findings}
        for metric in ("avg_read_io_cache_percent_30", "tagged_query_pct_30", "dashboard_queries_30", "job_queries_30"):
            self.assertIn(metric, metrics)

    def test_identical_evidence_merge_deduplicates(self):
        with tempfile.TemporaryDirectory() as td:
            a, b = Path(td) / "a.json", Path(td) / "b.json"
            p.save_evidence(a, self.evidence)
            p.save_evidence(b, self.evidence)
            merged = p.load_evidence_files([str(a), str(b)])
            self.assertEqual(len(merged), 1)
            self.assertEqual(len(merged[0].epochs), 2)

    def test_same_warehouse_id_in_different_workspaces_is_not_collapsed(self):
        other = copy.deepcopy(self.evidence[0])
        other.workspace_id = "9999999999999999"
        other.current_epoch_id = p.epoch_id(other.workspace_id, other.warehouse_id, other.current_epoch.epoch_start)
        for ep in other.epochs:
            ep.workspace_id = other.workspace_id
            ep.epoch_id = p.epoch_id(other.workspace_id, other.warehouse_id, ep.epoch_start)
            if ep.is_current:
                other.current_epoch_id = ep.epoch_id
        for row in other.daily_cost_365:
            row["workspace_id"] = other.workspace_id
        with tempfile.TemporaryDirectory() as td:
            a, b = Path(td) / "a.json", Path(td) / "b.json"
            p.save_evidence(a, self.evidence)
            p.save_evidence(b, [other])
            merged = p.load_evidence_files([str(a), str(b)])
            self.assertEqual(len(merged), 2)

    def test_mixed_currency_rejected(self):
        model = p.Modeler(self.t)
        rows = [
            {"usage_date": "2026-01-01", "list_cost": 1, "currency_code": "USD"},
            {"usage_date": "2026-01-02", "list_cost": 1, "currency_code": "EUR"},
        ]
        with self.assertRaises(ValueError):
            model.forecast_from_daily_cost(rows, "2025-01-01T00:00:00+00:00", "2026-10-01T00:00:00+00:00")

    def test_forecast_zero_fills_calendar_days(self):
        model = p.Modeler(self.t)
        rows = [
            {"usage_date": "2026-09-28", "list_cost": 100, "currency_code": "USD"},
            {"usage_date": "2026-09-30", "list_cost": 100, "currency_code": "USD"},
        ]
        out = model.forecast_from_daily_cost(rows, "2026-09-27T00:00:00+00:00", "2026-10-01T00:00:00+00:00")
        self.assertEqual(out["cost_history_days"], 4)
        self.assertEqual(out["active_cost_days"], 2)
        self.assertEqual(out["zero_cost_days_in_basis"], 2)
        self.assertAlmostEqual(out["forecast_cost_365"], 50 * 365, places=2)

    def _cost_rows(self, days, start="2026-09-01", cost=lambda i: 100.0):
        import datetime as dt
        d0 = dt.date.fromisoformat(start)
        return [{"usage_date": (d0 + dt.timedelta(days=i)).isoformat(), "list_cost": cost(i), "currency_code": "USD"} for i in range(days)]

    def test_cv_is_unknown_when_history_is_too_short(self):
        # 4 days < forecast_cv_min_days (7): the old code published CV 0.0 = "perfectly stable".
        model = p.Modeler(self.t)
        out = model.forecast_from_daily_cost(self._cost_rows(4, "2026-09-26"), "2026-09-26T00:00:00+00:00", "2026-10-01T00:00:00+00:00")
        self.assertIsNone(out["cost_cv_84"])

    def test_cv_is_unknown_without_any_cost_basis(self):
        model = p.Modeler(self.t)
        out = model.forecast_from_daily_cost([], "2026-10-01T00:00:00+00:00", "2026-10-01T00:00:00+00:00")
        self.assertIsNone(out["cost_cv_84"])

    def test_cv_is_unknown_when_mean_cost_is_zero(self):
        model = p.Modeler(self.t)
        out = model.forecast_from_daily_cost(self._cost_rows(30, cost=lambda i: 0.0), "2026-09-01T00:00:00+00:00", "2026-10-01T00:00:00+00:00")
        self.assertIsNone(out["cost_cv_84"])

    def test_cv_is_computed_when_history_is_sufficient(self):
        model = p.Modeler(self.t)
        out = model.forecast_from_daily_cost(self._cost_rows(30, cost=lambda i: 50.0 if i % 2 else 150.0), "2026-09-01T00:00:00+00:00", "2026-10-01T00:00:00+00:00")
        self.assertAlmostEqual(out["cost_cv_84"], 0.5, places=3)

    def test_unknown_cv_is_unknown_stability_not_perfect_stability(self):
        known = copy.deepcopy(self.evidence)
        unknown = copy.deepcopy(self.evidence)
        known[0].cost_cv_84 = 0.0
        unknown[0].cost_cv_84 = None
        c_known = self.candidate_of(known, "warehouse_size")
        c_unknown = self.candidate_of(unknown, "warehouse_size")
        self.assertEqual(c_known.confidence["components"]["stability"], 1.0)
        self.assertIsNone(c_unknown.confidence["components"]["stability"])
        # Unknown stability earns no credit, so confidence can only go down (some candidates sit at a cap).
        scores = [(self.candidate_of(known, k).confidence["score"], self.candidate_of(unknown, k).confidence["score"]) for k in p.ACTIVE_KNOBS]
        self.assertTrue(all(u <= k for k, u in scores))
        self.assertTrue(any(u < k for k, u in scores))
        fc = p._forecast_confidence(unknown[0], self.t)
        self.assertIsNone(fc.components["cost_stability"])
        self.assertLess(fc.score, p._forecast_confidence(known[0], self.t).score)

    def candidate_of(self, evidence, knob):
        return next(c for c in p.evaluate_warehouses(evidence, self.t)[0].candidates if c.knob == knob)

    def test_unknown_cv_survives_evidence_roundtrip_and_merge(self):
        self.evidence[0].cost_cv_84 = None
        with tempfile.TemporaryDirectory() as td:
            a, b = Path(td) / "a.json", Path(td) / "b.json"
            p.save_evidence(a, self.evidence)
            p.save_evidence(b, self.evidence)
            self.assertIsNone(p.load_evidence_files([str(a)])[0].cost_cv_84)
            merged = p.load_evidence_files([str(a), str(b)])
            self.assertIsNone(merged[0].cost_cv_84)

    def test_unknown_cv_packages_validate(self):
        self.evidence[0].cost_cv_84 = None
        pkgs = p.build_recommendation_packages(self.evidence, self.t)
        self.assertTrue(pkgs)

    def size_candidate_with_failure_rate(self, value, *, drop=False):
        ev = copy.deepcopy(self.evidence)
        q = ev[0].current_epoch.query
        if drop:
            q.pop("failure_cancel_pct_30", None)
        else:
            q["failure_cancel_pct_30"] = value
        return self.candidate_of(ev, "warehouse_size")

    def test_elevated_failure_cancel_rate_blocks_size_down(self):
        # Demo warehouse downsizes LARGE->MEDIUM at 0.2% failures; at 8% the "no pressure" evidence is not trustworthy.
        self.assertEqual(self.size_candidate_with_failure_rate(0.2).status, "TEST_REQUIRED")
        c = self.size_candidate_with_failure_rate(8.0)
        self.assertEqual(c.status, "NO_CHANGE")
        self.assertIn("FAILURE_CANCEL_RATE_HIGH", c.blockers)
        self.assertEqual(c.recommended_config, c.current_config)

    def test_failure_cancel_threshold_boundary_is_inclusive(self):
        limit = self.t.failure_cancel_pct_downsize_safe
        self.assertEqual(self.size_candidate_with_failure_rate(limit).status, "TEST_REQUIRED")
        self.assertIn("FAILURE_CANCEL_RATE_HIGH", self.size_candidate_with_failure_rate(limit + 0.01).blockers)

    def test_missing_failure_cancel_evidence_fails_closed_for_size_down(self):
        for c in (self.size_candidate_with_failure_rate(None, drop=True), self.size_candidate_with_failure_rate(None)):
            self.assertEqual(c.status, "NO_CHANGE")
            self.assertIn("FAILURE_CANCEL_EVIDENCE_UNAVAILABLE", c.blockers)

    def test_failure_cancel_rate_is_reported_as_a_size_finding(self):
        ev = copy.deepcopy(self.evidence)
        ev[0].current_epoch.query["failure_cancel_pct_30"] = 8.0
        analysis = p.Analyzer(self.t).analyze(ev[0])["warehouse_size"]
        finding = next(f for f in analysis.findings if f.metric == "failure_cancel_pct_30")
        self.assertEqual(finding.value, 8.0)
        self.assertEqual(finding.threshold, self.t.failure_cancel_pct_downsize_safe)

    def serverless_size_evidence(self, size, golden="serverless", **query):
        payload = json.loads((ROOT / "tests" / "golden" / f"{golden}.json").read_text())
        ev = payload["evidence"][0]
        for ep in ev["epochs"]:
            ep["config"]["warehouse_size"] = size
            ep["config"]["warehouse_type"] = "SERVERLESS"
        next(x for x in ev["epochs"] if x["is_current"])["query"].update(query)
        return [p.evidence_from_dict(ev)]

    def test_serverless_5x_large_downsize_is_not_priced_from_a_nominal_ratio(self):
        # 5X-Large has no published Serverless DBU/hour. Before: 5X->4X was priced at a guessed 0.5 ratio (~$9,450).
        c = self.candidate_of(self.serverless_size_evidence("5X_LARGE"), "warehouse_size")
        self.assertEqual(c.status, "TEST_REQUIRED")
        self.assertEqual(c.recommended_config, {"warehouse_size": "4X_LARGE"})
        self.assertIsNone(c.estimated_cost_365)
        self.assertIsNone(c.estimated_savings_365)
        self.assertIsNone(c.total_365d_savings_estimate)
        self.assertIsNone(c.dbu_365d_savings_range)
        self.assertEqual(c.metadata["estimate_blocker"], "SERVERLESS_DBU_RATE_NOT_PUBLISHED")
        self.assertNotIn("capacity_ratio_source", c.metadata)
        self.assertEqual(c.metadata["bundle_effect"]["type"], "requires_canary")

    def test_serverless_upsize_into_5x_large_is_not_priced_from_a_nominal_ratio(self):
        ev = self.serverless_size_evidence(
            "4X_LARGE", golden="spill-constrained",
        )
        c = self.candidate_of(ev, "warehouse_size")
        self.assertEqual(c.recommended_config, {"warehouse_size": "5X_LARGE"})
        self.assertIsNone(c.estimated_savings_365)
        self.assertEqual(c.metadata["estimate_blocker"], "SERVERLESS_DBU_RATE_NOT_PUBLISHED")

    def test_serverless_sizes_with_published_rates_are_still_priced(self):
        c = self.candidate_of(self.serverless_size_evidence("4X_LARGE"), "warehouse_size")
        self.assertEqual(c.recommended_config, {"warehouse_size": "3X_LARGE"})
        self.assertGreater(c.estimated_savings_365, 0)
        self.assertEqual(c.metadata["capacity_ratio_source"], "DATABRICKS_PUBLISHED_SQL_SERVERLESS_DBU_PER_HOUR_RATIO")
        self.assertNotIn("estimate_blocker", c.metadata)

    def test_unpriceable_serverless_size_step_withholds_scenario_economics(self):
        ev = self.serverless_size_evidence("5X_LARGE")
        w = p.evaluate_warehouses(ev, self.t)[0]
        scenario = next(s for s in w.scenarios if any(st.knob == "warehouse_size" for st in s.steps))
        self.assertEqual(scenario.portfolio_estimate.estimate_status, "REQUIRES_CANARY")
        self.assertIsNone(scenario.portfolio_estimate.total_365d_savings_range)
        self.assertIsNone(scenario.portfolio_estimate.dbu_365d_savings_range)
        self.assertTrue(p.build_recommendation_packages(ev, self.t))

    def test_withheld_serverless_size_candidate_is_published_as_unavailable(self):
        pkg = p.build_recommendation_packages(self.serverless_size_evidence("5X_LARGE"), self.t)[0]
        size = next(c for c in pkg.candidates if c.knob.value == "warehouse_size")
        self.assertEqual(size.status.value, "TEST_REQUIRED")
        self.assertEqual(size.independent_economics.estimation_status.value, "UNAVAILABLE")
        self.assertIsNone(size.independent_economics.savings_365d.total)

    def with_cost_history(self, days, golden=None):
        if golden:
            payload = json.loads((ROOT / "tests" / "golden" / f"{golden}.json").read_text())
            ev = [p.evidence_from_dict(payload["evidence"][0])]
        else:
            ev = copy.deepcopy(self.evidence)
        ev[0].cost_history_days = days
        return ev

    def test_migration_needs_the_cost_history_floor(self):
        floor = self.t.min_economics_cost_history_days
        self.assertEqual(floor, 7)
        short = self.candidate_of(self.with_cost_history(floor - 1), "warehouse_type")
        self.assertEqual(short.status, "INSUFFICIENT_EVIDENCE")
        self.assertEqual(short.blockers, ["COST_HISTORY_INSUFFICIENT"])
        self.assertEqual(short.recommended_config, short.current_config)
        self.assertIsNone(short.estimated_savings_365)
        ok = self.candidate_of(self.with_cost_history(floor), "warehouse_type")
        self.assertEqual(ok.status, "TEST_REQUIRED")
        self.assertIsNotNone(ok.estimated_savings_365)

    def test_short_cost_history_keeps_the_decision_but_withholds_economics(self):
        # Queue pressure on a 6-day-old epoch: the performance decision stands (it does not rest on cost), the number does not.
        short = self.candidate_of(self.with_cost_history(6, "concurrency-constrained"), "max_clusters")
        full = self.candidate_of(self.with_cost_history(7, "concurrency-constrained"), "max_clusters")
        self.assertEqual(short.status, full.status)
        self.assertEqual(short.recommended_config, full.recommended_config)
        self.assertIsNone(short.estimated_savings_365)
        self.assertIsNone(short.estimated_cost_range_365)
        self.assertEqual(short.metadata["estimate_blocker"], "COST_HISTORY_INSUFFICIENT")
        self.assertEqual(short.metadata["bundle_effect"]["type"], "estimate_withheld")
        self.assertNotIn("estimate_blocker", full.metadata)

    def test_short_cost_history_does_not_touch_no_change_candidates(self):
        for c in p.evaluate_warehouses(self.with_cost_history(3, "newly-changed-config"), self.t)[0].candidates:
            if c.status == "NO_CHANGE":
                self.assertNotIn("estimate_blocker", c.metadata)

    def test_short_cost_history_withholds_scenario_economics_and_packages_validate(self):
        ev = self.with_cost_history(6, "concurrency-constrained")
        w = p.evaluate_warehouses(ev, self.t)[0]
        plan_a = next(s for s in w.scenarios if s.steps and any(st.knob == "max_clusters" for st in s.steps))
        self.assertEqual(plan_a.portfolio_estimate.estimate_status, "REQUIRES_COST_HISTORY")
        self.assertIsNone(plan_a.portfolio_estimate.total_365d_savings_range)
        pkg = p.build_recommendation_packages(ev, self.t)[0]
        size = next(c for c in pkg.candidates if c.knob.value == "max_clusters")
        self.assertEqual(size.independent_economics.estimation_status.value, "UNAVAILABLE")

    # --- M7: billing window boundary and completeness watermark -------------------------------------------------
    ANALYSIS = "2026-10-01T12:00:00+00:00"

    def forecast(self, rows, watermark="UNSET", start="2026-09-01T00:00:00+00:00"):
        kwargs = {} if watermark == "UNSET" else {"billing_watermark": watermark}
        return p.Modeler(self.t).forecast_from_daily_cost(rows, start, self.ANALYSIS, **kwargs)

    def test_billing_sql_uses_an_exact_half_open_window_that_excludes_today(self):
        for name in ("EPOCH_DAILY_COST_SQL", "BILLING_FEATURE_SQL", "EPOCH_BILLING_ATTRIBUTION_SQL"):
            sql = getattr(p, name)
            self.assertIn("usage_date>=DATE_SUB(CURRENT_DATE(),365)", sql, name)
            self.assertIn("usage_date<CURRENT_DATE()", sql, name)

    def test_billing_cost_sql_keeps_corrections_in_the_sum(self):
        # RETRACTION (negative) and RESTATEMENT rows net out original usage; filtering to ORIGINAL would overstate cost.
        for name in ("EPOCH_DAILY_COST_SQL", "BILLING_FEATURE_SQL", "EPOCH_BILLING_ATTRIBUTION_SQL"):
            self.assertNotIn("record_type", getattr(p, name), name)

    def test_billing_watermark_sql_is_read_only_and_bounded(self):
        sql = p.BILLING_WATERMARK_SQL
        self.assertIn("MAX(usage_end_time)", sql)
        self.assertIn("system.billing.usage", sql)
        self.assertIn("usage_date>=DATE_SUB(CURRENT_DATE()", sql)
        self.assertNotRegex(sql.upper(), r"\b(INSERT|UPDATE|DELETE|DROP|ALTER|CREATE)\b")

    def test_day_not_yet_covered_by_the_watermark_is_excluded_from_the_basis(self):
        rows = self._cost_rows(30, "2026-09-01")  # 09-01 .. 09-30, $100/day; 09-30 is only partly ingested
        rows[-1]["list_cost"] = 20.0
        out = self.forecast(rows, watermark="2026-09-30T20:00:00+00:00")
        self.assertEqual(out["billing_complete_through"], "2026-09-29")
        self.assertEqual(out["billing_completeness"], "COVERED_BY_WATERMARK")
        self.assertEqual(out["excluded_incomplete_days"], 1)
        self.assertEqual(out["cost_history_days"], 29)
        self.assertAlmostEqual(out["forecast_cost_365"], 100 * 365, places=2)
        self.assertAlmostEqual(out["current_cost_365"], 29 * 100, places=2)

    def test_watermark_exactly_at_midnight_completes_the_previous_day(self):
        rows = self._cost_rows(30, "2026-09-01")
        self.assertEqual(self.forecast(rows, watermark="2026-10-01T00:00:00+00:00")["billing_complete_through"], "2026-09-30")
        self.assertEqual(self.forecast(rows, watermark="2026-09-30T23:00:00+00:00")["billing_complete_through"], "2026-09-29")

    def test_watermark_never_extends_the_basis_past_yesterday(self):
        out = self.forecast(self._cost_rows(30, "2026-09-01"), watermark="2026-10-02T06:00:00+00:00")
        self.assertEqual(out["billing_complete_through"], "2026-09-30")
        self.assertEqual(out["excluded_incomplete_days"], 0)

    def test_missing_watermark_is_labeled_unverified_not_silently_complete(self):
        out = self.forecast(self._cost_rows(30, "2026-09-01"))
        self.assertEqual(out["billing_completeness"], "UNVERIFIED_ASSUMED_COMPLETE_THROUGH_PRIOR_DAY")
        self.assertIsNone(out["billing_watermark"])
        self.assertEqual(out["billing_complete_through"], "2026-09-30")
        also_none = self.forecast(self._cost_rows(30, "2026-09-01"), watermark=None)
        self.assertEqual(also_none["billing_completeness"], "UNVERIFIED_ASSUMED_COMPLETE_THROUGH_PRIOR_DAY")

    def test_actual_cost_uses_the_same_half_open_365_day_window(self):
        rows = [
            {"usage_date": "2025-09-30", "list_cost": 1000.0, "currency_code": "USD"},  # 366 days back: outside
            {"usage_date": "2025-10-01", "list_cost": 10.0, "currency_code": "USD"},   # 365 days back: inside
            {"usage_date": "2026-09-30", "list_cost": 10.0, "currency_code": "USD"},   # yesterday: inside
            {"usage_date": "2026-10-01", "list_cost": 500.0, "currency_code": "USD"},  # today (incomplete): outside
        ]
        out = self.forecast(rows, start="2025-10-01T00:00:00+00:00")
        self.assertEqual(out["current_cost_365"], 20.0)

    # --- G2: usage with no list_prices row is unavailable cost, never silently dropped ------------------------
    def test_cost_sql_keeps_unpriced_usage_and_one_price_per_usage_record(self):
        sql = p.EPOCH_DAILY_COST_SQL
        self.assertIn("LEFT JOIN system.billing.list_prices", sql)
        self.assertNotRegex(sql, r"\)\s*JOIN system\.billing\.list_prices|\bu\s+JOIN system\.billing\.list_prices")
        self.assertIn("unpriced_dbu_quantity", sql)
        # Overlapping price rows must not double count one usage record.
        self.assertIn("PARTITION BY u.record_id", sql)

    def _rows_with_unpriced(self, unpriced_dates):
        rows = [{"usage_date": f"2026-09-{d:02d}", "list_cost": 100.0, "dbu_quantity": 200.0, "unpriced_dbu_quantity": 0.0, "currency_code": "USD"}
                for d in range(1, 31)]
        rows += [{"usage_date": d, "list_cost": None, "dbu_quantity": 150.0, "unpriced_dbu_quantity": 150.0, "currency_code": None}
                 for d in unpriced_dates]
        return rows

    def test_unpriced_usage_in_forecast_basis_makes_forward_and_actual_cost_unknown(self):
        out = self.forecast(self._rows_with_unpriced(["2026-09-20", "2026-09-21"]))
        self.assertIsNone(out["forecast_cost_365"])
        self.assertIsNone(out["current_cost_365"])
        self.assertEqual(out["unpriced_dbu_quantity_365"], 300.0)
        self.assertEqual(out["unpriced_dbu_quantity_in_basis"], 300.0)
        self.assertEqual(out["dbu_pricing_status"], "UNPRICED_USAGE")
        self.assertEqual(out["currency"], "USD")  # unpriced rows carry no currency and do not count as mixed

    def test_unpriced_usage_before_current_epoch_only_nulls_the_trailing_actual(self):
        rows = self._rows_with_unpriced(["2026-08-15"])
        out = self.forecast(rows)
        self.assertIsNone(out["current_cost_365"])
        self.assertIsNotNone(out["forecast_cost_365"])
        self.assertEqual(out["unpriced_dbu_quantity_in_basis"], 0.0)
        self.assertEqual(out["dbu_pricing_status"], "UNPRICED_USAGE")

    def test_fully_priced_usage_is_complete(self):
        out = self.forecast(self._rows_with_unpriced([]))
        self.assertEqual(out["dbu_pricing_status"], "COMPLETE")
        self.assertEqual(out["unpriced_dbu_quantity_365"], 0.0)
        self.assertAlmostEqual(out["current_cost_365"], 3000.0)

    def test_collector_marks_unpriced_epoch_cost_unknown(self):
        class Client:
            def query(self_inner, statement, *, label="collector_sql"):
                if label == "config_epochs":
                    return [{"workspace_id": "1", "warehouse_id": "w", "account_id": "a", "warehouse_name": "W", "warehouse_type": "PRO",
                             "warehouse_channel": "CHANNEL_NAME_CURRENT", "warehouse_size": "MEDIUM", "min_clusters": 1, "max_clusters": 2,
                             "auto_stop_minutes": 10, "epoch_start": "2026-09-01T00:00:00+00:00", "epoch_end": None, "is_current": True, "tags": {}}]
                if label == "billing_daily_cost":
                    return [dict(r, workspace_id="1", warehouse_id="w", epoch_start="2026-09-01T00:00:00+00:00")
                            for r in self._rows_with_unpriced(["2026-09-20"])]
                return []
        e = p.extract_evidence(Client(), self.t, "us-east-1")[0]
        self.assertIsNone(e.baseline_forecast_cost_365)
        self.assertIsNone(e.current_cost_365)
        self.assertIsNone(e.current_epoch.cost["list_cost"])
        self.assertEqual(e.current_epoch.cost["unpriced_dbu_quantity"], 150.0)
        self.assertEqual(e.current_epoch.cost["dbu_quantity"], 6150.0)

    def test_unpriced_usage_in_basis_leaves_trend_and_stability_unknown(self):
        # A run of unpriced days is not a cost drop: trend, run rates and CV must not read it as $0 days.
        rows = [{"usage_date": f"2026-{m:02d}-{d:02d}", "list_cost": 100.0, "dbu_quantity": 200.0, "unpriced_dbu_quantity": 0.0, "currency_code": "USD"}
                for m, days in ((5, 31), (6, 30), (7, 31), (8, 31)) for d in range(1, days + 1)]
        rows += [{"usage_date": f"2026-09-{d:02d}", "list_cost": None, "dbu_quantity": 200.0, "unpriced_dbu_quantity": 200.0, "currency_code": None}
                 for d in range(1, 31)]
        out = self.forecast(rows, start="2026-05-01T00:00:00+00:00")
        self.assertIsNone(out["forecast_cost_365"])
        self.assertIsNone(out["trend_ratio"])
        self.assertEqual(out["trend_adjustment_factor"], 1.0)
        self.assertIsNone(out["run_rate_28"])
        self.assertIsNone(out["run_rate_84"])
        self.assertIsNone(out["cost_cv_84"])

    def test_cost_sql_treats_a_null_list_price_as_unpriced_and_breaks_price_ties(self):
        sql = p.EPOCH_DAILY_COST_SQL
        self.assertIn("p.pricing.effective_list.default IS NULL", sql)
        self.assertIn("ORDER BY p.price_start_time DESC NULLS LAST, p.currency_code", sql)

    def test_evidence_with_unpriced_rows_merges_with_itself(self):
        payload = json.loads((ROOT / "tests" / "golden" / "c32-unpriced-sku.json").read_text())
        with tempfile.TemporaryDirectory() as td:
            a, b = Path(td) / "a.json", Path(td) / "b.json"
            for path in (a, b):
                path.write_text(json.dumps(payload["evidence"]))
            merged = p.load_evidence_files([str(a), str(b)])
        self.assertEqual(len(merged), 1)
        self.assertEqual(len(merged[0].daily_cost_365), len(payload["evidence"][0]["daily_cost_365"]))

    def test_unknown_cost_survives_evidence_round_trip(self):
        e = copy.deepcopy(self.evidence[0])
        e.current_cost_365 = None
        e.baseline_forecast_cost_365 = None
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "e.json"
            p.save_evidence(path, [e])
            back = p.load_evidence_files([str(path)])[0]
        self.assertIsNone(back.current_cost_365)
        self.assertIsNone(back.baseline_forecast_cost_365)

    def test_unknown_dbu_baseline_withholds_estimates_and_migration(self):
        e = copy.deepcopy(self.evidence[0])
        e.baseline_forecast_cost_365 = None
        e.current_cost_365 = None
        w = p.evaluate_warehouses([e], self.t)[0]
        by = {c.knob: c for c in w.candidates}
        self.assertEqual(by["warehouse_type"].status, "INSUFFICIENT_EVIDENCE")
        self.assertIn("DBU_COST_UNAVAILABLE", by["warehouse_type"].blockers)
        for c in w.candidates:
            self.assertIsNone(c.dbu_365d_baseline_cost, c.knob)
            if c.status not in {"NO_CHANGE"}:
                self.assertIsNone(c.estimated_savings_range_365, c.knob)
        pkg = p.build_recommendation_packages([e], self.t)[0].to_dict()
        p.validate_locked_package_schema(pkg)
        self.assertIsNone(pkg["cost_baseline"]["forward_365d"]["dbu_cost"])
        self.assertEqual(pkg["cost_baseline"]["forward_365d"]["availability"], "PARTIAL")

    def test_collector_passes_the_billing_watermark_to_the_forecast(self):
        seen = {}
        real = p.Modeler.forecast_from_daily_cost
        def spy(model, rows, start, analysis_time, **kw):
            seen.update(kw)
            return real(model, rows, start, analysis_time, **kw)
        class Client:
            def query(self_inner, statement, *, label="collector_sql"):
                if label == "billing_watermark":
                    return [{"billing_watermark": "2026-09-30T20:00:00+00:00"}]
                if label == "config_epochs":
                    return [{"workspace_id": "1", "warehouse_id": "w", "account_id": "a", "warehouse_name": "W", "warehouse_type": "PRO",
                             "warehouse_channel": "CHANNEL_NAME_CURRENT", "warehouse_size": "MEDIUM", "min_clusters": 1, "max_clusters": 2,
                             "auto_stop_minutes": 10, "epoch_start": "2026-09-01T00:00:00+00:00", "epoch_end": None, "is_current": True, "tags": {}}]
                return []
        with patch.object(p.Modeler, "forecast_from_daily_cost", spy):
            p.extract_evidence(Client(), self.t, "us-east-1")
        self.assertEqual(p.parse_datetime(seen["billing_watermark"]).isoformat(), "2026-09-30T20:00:00+00:00")

    def auto_stop_only(self, idle_multiple):
        """Idle-expensive fixture where auto-stop is the only plan step (no size/min/max change)."""
        import dataclasses
        payload = json.loads((ROOT / "tests" / "golden" / "idle-expensive-warehouse.json").read_text())
        ev = p.evidence_from_dict(payload["evidence"][0])
        ep = ev.current_epoch
        ep.config = dataclasses.replace(ep.config, min_clusters=1)
        ep.query.update(queued_query_pct_30=0.6, queued_queries_30=72, p95_wait_capacity_ms_when_queued_30=1000)
        for k in list(ep.gaps):
            if k.startswith("saved_idle_minutes_if_"):
                ep.gaps[k] = ep.events["cluster_minutes"] * idle_multiple
        w = p.evaluate_warehouses([ev], self.t)[0]
        cand = next(c for c in w.candidates if c.knob == "auto_stop_minutes")
        plan = next(s for s in w.scenarios if [st.knob for st in s.steps] == ["auto_stop_minutes"])
        return cand, plan

    def test_single_step_auto_stop_scenario_matches_the_capped_candidate(self):
        # Idle gaps worth 90% of all cluster-minutes: the candidate caps removable time at auto_stop_savings_cap (60%).
        # Before: the stateful bundle used the uncapped idle minutes, so the one-step plan claimed more than its own step.
        for multiple in (0.2, 0.9, 2.0):
            cand, plan = self.auto_stop_only(multiple)
            self.assertAlmostEqual(plan.portfolio_estimate.dbu_365d_savings_range.base, cand.dbu_365d_savings_estimate, places=2, msg=multiple)
            self.assertAlmostEqual(plan.portfolio_estimate.dbu_365d_savings_range.high, cand.dbu_365d_savings_range.high, places=2, msg=multiple)
        capped, _ = self.auto_stop_only(2.0)
        self.assertEqual(capped.metadata["raw_savings_fraction"], self.t.auto_stop_savings_cap)

    def test_guardrail_changes_with_no_direct_savings_are_labeled_no_delta_modeled(self):
        # Demo max-cluster tightening (4 -> 3) is a guardrail: it carries no priced delta, so it is not a "modeled range".
        pkg = self.packages[0]
        mx = next(c for c in pkg.candidates if c.knob.value == "max_clusters")
        self.assertEqual(mx.status.value, "ACTIONABLE")
        self.assertEqual(mx.independent_economics.estimation_status.value, "NO_DELTA_MODELED")
        self.assertEqual(mx.independent_economics.savings_365d.total.base, 0)

    def test_priced_changes_keep_their_modeled_label(self):
        up = self.increase_candidate_package()
        self.assertEqual(up.status.value, "ACTIONABLE")
        self.assertEqual(up.independent_economics.estimation_status.value, "MODELED_RANGE")
        for c in self.packages[0].candidates:
            if c.knob.value in {"auto_stop_minutes", "min_clusters", "warehouse_size"} and c.status.value in {"ACTIONABLE", "TEST_REQUIRED"}:
                self.assertNotEqual(c.independent_economics.estimation_status.value, "NO_DELTA_MODELED", c.knob.value)

    def increase_candidate_package(self):
        payload = json.loads((ROOT / "tests" / "golden" / "concurrency-constrained.json").read_text())
        pkg = p.build_recommendation_packages([p.evidence_from_dict(payload["evidence"][0])], self.t)[0]
        return next(c for c in pkg.candidates if c.knob.value == "max_clusters")

    def test_planner_ready_for_hitl_maps_explicitly_to_contract_ready(self):
        payload = json.loads((ROOT / "tests" / "golden" / "concurrency-constrained.json").read_text())
        ev = [p.evidence_from_dict(payload["evidence"][0])]
        w = p.evaluate_warehouses(ev, self.t)[0]
        internal = {st.knob: st.execution_state for s in w.scenarios for st in s.steps}
        self.assertEqual(internal["max_clusters"], "READY_FOR_HITL")
        pkg = p.build_recommendation_packages(ev, self.t)[0]
        states = {st.knob.value: st.execution_state.value for sc in pkg.scenarios for st in sc.execution_plan if getattr(st, "knob", None)}
        self.assertEqual(states["max_clusters"], "READY")

    def test_unknown_planner_execution_state_fails_loudly_instead_of_becoming_ready(self):
        payload = json.loads((ROOT / "tests" / "golden" / "concurrency-constrained.json").read_text())
        ev = [p.evidence_from_dict(payload["evidence"][0])]
        real = p.ScenarioPlanner.scenarios
        def corrupt(planner, e, candidates):
            scenarios, recommended = real(planner, e, candidates)
            scenarios[0].steps[0].execution_state = "NOT_A_STATE"
            return scenarios, recommended
        with patch.object(p.ScenarioPlanner, "scenarios", corrupt):
            with self.assertRaises(ValueError) as ctx:
                p.build_recommendation_packages(ev, self.t)
        self.assertIn("NOT_A_STATE", str(ctx.exception))

    def envelope_cases(self):
        """Serverless golden: Plan A is priced, Plan B (Pro migration) is UNAVAILABLE."""
        payload = json.loads((ROOT / "tests" / "golden" / "serverless.json").read_text())
        e = p.evidence_from_dict(payload["evidence"][0])
        pkg = p.build_recommendation_packages([e], self.t)[0]
        priced = next(s for s in pkg.scenarios if s.portfolio_economics.savings_365d.total is not None)
        unpriced = next(s for s in pkg.scenarios if s.portfolio_economics.savings_365d.total is None)
        return e, pkg.scenarios, priced, unpriced

    def reasons(self, e, scenarios, recommended):
        econ = p._package_economics(e, scenarios, recommended.scenario_id if recommended else None)
        return {x.scenario_id: x.reason_code for x in econ.excluded_scenarios}, econ

    def test_unpriced_alternative_is_not_blamed_on_the_recommended_scenario(self):
        e, scenarios, priced, unpriced = self.envelope_cases()
        why, econ = self.reasons(e, scenarios, priced)
        self.assertEqual(why, {unpriced.scenario_id: "SCENARIO_ECONOMICS_UNAVAILABLE"})
        self.assertEqual(econ.included_scenario_ids, [priced.scenario_id])

    def test_envelope_stays_empty_when_the_recommended_scenario_is_unpriced(self):
        e, scenarios, priced, unpriced = self.envelope_cases()
        why, econ = self.reasons(e, scenarios, unpriced)
        self.assertEqual(why, {
            unpriced.scenario_id: "SCENARIO_ECONOMICS_UNAVAILABLE",
            priced.scenario_id: "RECOMMENDED_SCENARIO_ECONOMICS_UNAVAILABLE",
        })
        self.assertIsNone(econ.estimated_savings_365d)
        self.assertEqual(econ.included_scenario_ids, [])

    def test_priced_scenario_with_no_recommendation_gets_its_own_reason(self):
        e, scenarios, priced, unpriced = self.envelope_cases()
        why, econ = self.reasons(e, scenarios, None)
        self.assertEqual(why[priced.scenario_id], "NO_RECOMMENDED_SCENARIO")
        self.assertEqual(why[unpriced.scenario_id], "SCENARIO_ECONOMICS_UNAVAILABLE")

    def test_watermark_sql_looks_back_far_enough_to_catch_a_stalled_feed_and_is_sql_scoped(self):
        sql = p.BILLING_WATERMARK_SQL
        self.assertIn("usage_date>=DATE_SUB(CURRENT_DATE(),30)", sql)
        self.assertIn("billing_origin_product='SQL'", sql)

    def test_replayed_short_history_cv_of_zero_is_unknown(self):
        # Evidence written before the CV fix stored 0.0 for short histories; replay must not grant full stability credit.
        payload = json.loads((ROOT / "tests" / "golden" / "serverless.json").read_text())["evidence"][0]
        payload["cost_history_days"], payload["cost_cv_84"] = 3, 0.0
        self.assertIsNone(p.evidence_from_dict(payload).cost_cv_84)
        payload["cost_history_days"], payload["cost_cv_84"] = 30, 0.0
        self.assertEqual(p.evidence_from_dict(payload).cost_cv_84, 0.0)

    def test_demo_cv_is_unknown_for_a_warehouse_with_too_little_history(self):
        for e in p.synthetic_demo_evidence(15):
            if e.cost_history_days < self.t.forecast_cv_min_days:
                self.assertIsNone(e.cost_cv_84, e.warehouse_name)

    def test_gap_collector_handles_overlapping_queries_and_policy_targets(self):
        sql = p.build_epoch_gap_metrics_sql(self.t)
        self.assertIn("MAX(x.end_time) OVER", sql)
        for target in {self.t.serverless_auto_stop_target_minutes, self.t.low_activity_auto_stop_target_minutes, self.t.medium_activity_auto_stop_target_minutes, self.t.high_activity_auto_stop_target_minutes}:
            self.assertIn(f"saved_idle_minutes_if_{target}", sql)

    def test_profile_host_and_cli_oauth_flow(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = Path(td) / ".databrickscfg"
            cfg.write_text("[WHOPT]\nhost = https://dbc.example.cloud.databricks.com\n")
            self.assertEqual(p._load_profile_host("WHOPT", str(cfg)), "dbc.example.cloud.databricks.com")
            class R:
                stdout = '{"access_token":"short-lived","expiry":"2026-10-01T20:00:00Z"}'
            with patch("publisher.subprocess.run", return_value=R()) as run:
                auth = p.resolve_cli_oauth("WHOPT", "abc123", str(cfg))
                self.assertEqual(auth.http_path, "/sql/1.0/warehouses/abc123")
                self.assertEqual(auth.access_token, "short-lived")
                run.assert_called_once()
                call = run.call_args
                self.assertEqual(call.kwargs["env"]["DATABRICKS_CONFIG_FILE"], str(cfg))
                self.assertNotIn("DATABRICKS_TOKEN", call.kwargs["env"])

    def test_all_purpose_cluster_execution_path_and_validation(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = Path(td) / ".databrickscfg"
            cfg.write_text("[WHOPT]\nhost = https://dbc.example.cloud.databricks.com\n")
            class R:
                stdout = '{"access_token":"short-lived","expiry":"2026-10-01T20:00:00Z"}'
            with patch("publisher.subprocess.run", return_value=R()), patch("publisher._validate_all_purpose_cluster", return_value={"cluster_source":"UI"}), patch("publisher._current_workspace_id", return_value="123456789"):
                auth = p.resolve_cli_oauth("WHOPT", "cluster", "1234-567890-demo123", str(cfg))
                self.assertEqual(auth.http_path, "sql/protocolv1/o/123456789/1234-567890-demo123")
                self.assertEqual(auth.execution_kind, "cluster")

    def test_compatibility_api_b_and_fallback_a(self):
        auth=p.AuthContext("WHOPT","dbc.example.cloud.databricks.com","/sql/1.0/warehouses/x","tok",None,"warehouse","x","123")
        payload={"enable_serverless_compute":True,"enabled_warehouse_types":[{"warehouse_type":"PRO","enabled":True}],"data_access_config":[],"security_policy":"NONE"}
        with patch("publisher._api_get", return_value=payload):
            ev=p.collect_workspace_compatibility_evidence(auth,"auto")
            self.assertEqual(ev["mode"],"B_READ_ONLY_API")
            self.assertTrue(ev["enable_serverless_compute"])
            self.assertFalse(ev["external_legacy_hive_metastore_detected"])
        with patch("publisher._api_get", side_effect=RuntimeError("denied")):
            ev=p.collect_workspace_compatibility_evidence(auth,"auto")
            self.assertEqual(ev["mode"],"A_HITL_FALLBACK")

    def test_candidate_and_recommended_plan_views(self):
        from contracts import workbook_contract as wc
        projection = wc.project_workbook(self.packages)
        row = projection.warehouse_portfolio[0]
        self.assertNotIn("Photon", row["candidate_view"])
        self.assertIn("Size", row["candidate_view"])
        self.assertIn("Type", row["candidate_view"])
        self.assertIn("warehouse_size", row["recommended_execution_plan"])
        self.assertIn("READY_FOR_CANARY", row["recommended_execution_plan"])
        self.assertIn("remeasure 7d", row["recommended_execution_plan"])

    def test_logging_file_is_created(self):
        with tempfile.TemporaryDirectory() as td:
            log=p.setup_logging(Path(td),"INFO")
            p.LOGGER.info("test.event value=%s","ok")
            self.assertTrue(log.exists())
            self.assertIn("test.event",log.read_text())

    def test_pat_profile_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = Path(td) / ".databrickscfg"
            cfg.write_text("[PAT]\nhost = https://dbc.example.cloud.databricks.com\ntoken = secret\n")
            with self.assertRaises(SystemExit):
                p.resolve_cli_oauth("PAT", "abc123", str(cfg))

    def test_same_evidence_produces_identical_locked_json_and_workbook_projection(self):
        from contracts import workbook_contract as wc
        ea = p.demo_evidence(); eb = p.demo_evidence()
        pa = p.build_recommendation_packages(ea, self.t)
        pb = p.build_recommendation_packages(eb, self.t)
        self.assertEqual([x.to_dict() for x in pa], [x.to_dict() for x in pb])
        self.assertEqual(wc.project_workbook(pa), wc.project_workbook(pb))
    def test_auto_stop_blocks_without_gap_or_event_evidence(self):
        ev = p.demo_evidence()
        ev[0].current_epoch.gaps = {}
        ev[0].current_epoch.events["cluster_minutes"] = 0
        result = p.evaluate_warehouses(ev, self.t)[0]
        auto = next(c for c in result.candidates if c.knob == "auto_stop_minutes")
        self.assertEqual(auto.status, "BLOCKED")
        self.assertIn("INSUFFICIENT_IDLE_GAP_SAMPLE", auto.blockers)
        self.assertIn("INSUFFICIENT_EVENT_RESOURCE_EVIDENCE", auto.blockers)

    def test_manifest_carries_locked_contract_and_workbook_metadata(self):
        with tempfile.TemporaryDirectory() as td:
            packages = self.packages
            files = p.package_locked_contract(Path(td), packages, self.evidence)
            manifest = json.loads(Path(files["manifest"]).read_text())
            self.assertEqual(manifest["publisher_version"], p.PUBLISHER_VERSION)
            self.assertEqual(manifest["reporter_contract_version"], "1.2.0")
            self.assertEqual(manifest["workbook_contract_version"], "1.3.0")
            self.assertEqual(manifest["workbook_status"], "GENERATED")
            self.assertEqual(manifest["package_count"], 1)
            self.assertEqual(manifest["workbook_file"], "workbook.xlsx")
    def test_reviewer_rejects_cross_workspace_candidate(self):
        candidates = list(self.w.candidates)
        bad = copy.deepcopy(candidates[0])
        bad.workspace_id = "wrong"
        candidates[0] = bad
        with self.assertRaises(ValueError):
            p.Reviewer(self.t).review(self.evidence[0], candidates)


    def test_photon_inference_uses_coverage_freshness_and_signed_dbu(self):
        row={"photon_true_dbu":950.0,"photon_false_dbu":50.0,"photon_null_dbu":0.0,"latest_usage_time":"2026-09-30T12:00:00+00:00"}
        obs=p.infer_photon_observation(row,"2026-10-01T12:00:00+00:00",self.t)
        self.assertEqual(obs.status,"ENABLED")
        self.assertEqual(obs.enabled_usage_pct,95.0)
        row["photon_null_dbu"]=500.0
        obs=p.infer_photon_observation(row,"2026-10-01T12:00:00+00:00",self.t)
        self.assertEqual(obs.status,"UNKNOWN")

    def test_owner_query_source_and_scenario_confidence_are_reported(self):
        self.assertEqual(self.w.evidence.attribution.owner,"sales-bi@databricks.example")
        self.assertEqual(self.w.query_source_profile["primary_category"],"DASHBOARD")
        self.assertEqual(self.w.query_source_profile["categories"]["POWER_BI"],4200)
        for sc in self.w.scenarios:
            self.assertIn("score",sc.confidence)
            for st in sc.steps:
                self.assertEqual(st.confidence["score"],self.by_knob[st.knob].confidence["score"])

    def test_normalized_domain_contracts_are_dataclasses(self):
        from dataclasses import is_dataclass
        for cls in (p.WarehouseConfig,p.EpochEvidence,p.PhotonObservation,p.QuerySourceProfile,p.WarehouseAttribution,p.AnalysisResult,p.OptimizationDecision,p.CandidateRecommendation,p.PlanStep,p.Scenario,p.WarehouseOptimization):
            self.assertTrue(is_dataclass(cls),cls.__name__)

    def test_pro_size_estimator_does_not_reuse_serverless_dbu_curve(self):
        c=self.by_knob["warehouse_size"]
        self.assertNotEqual(c.metadata.get("capacity_ratio_source"),"DATABRICKS_PUBLISHED_SQL_SERVERLESS_DBU_PER_HOUR_RATIO")
        self.assertIn(c.metadata.get("capacity_ratio_source"),{
            "NOMINAL_SIZE_CAPACITY_RATIO",
            "WORKSPACE_SYSTEM_TABLE_EMPIRICAL_PRO_CLASSIC_DBU_PER_CLUSTER_HOUR_RATIO",
        })


class GoldenScenarioTests(unittest.TestCase):
    def setUp(self):
        self.t=p.Thresholds(); self.golden=ROOT/"tests"/"golden"

    def load(self,name):
        payload=json.loads((self.golden/f"{name}.json").read_text())
        ev=[p.evidence_from_dict(x) for x in payload["evidence"]]
        return payload,p.evaluate_warehouses(ev,self.t)

    def warehouse(self,report,warehouse_id):
        return next(w for w in report if w.warehouse_id==warehouse_id)

    def test_all_golden_files_parse_and_are_deterministic(self):
        files=sorted(self.golden.glob("*.json")); self.assertGreaterEqual(len(files),10)
        for f in files:
            payload=json.loads(f.read_text()); ev=[p.evidence_from_dict(x) for x in payload["evidence"]]
            a=p.evaluate_warehouses(copy.deepcopy(ev),self.t); b=p.evaluate_warehouses(copy.deepcopy(ev),self.t)
            self.assertEqual(p.canonical_json(a),p.canonical_json(b),f.name)

    def test_idle_expensive(self):
        _,r=self.load("idle-expensive-warehouse"); w=self.warehouse(r,"wh-demo-001"); by={c.knob:c for c in w.candidates}
        self.assertEqual(by["min_clusters"].status,"TEST_REQUIRED"); self.assertEqual(by["auto_stop_minutes"].status,"ACTIONABLE")
        self.assertGreater(w.aws_365d_baseline_cost,0)

    def test_spill_constrained(self):
        _,r=self.load("spill-constrained"); c=next(c for c in r[0].candidates if c.knob=="warehouse_size")
        self.assertEqual(c.status,"ACTIONABLE"); self.assertEqual(c.recommended_config["warehouse_size"],"X_LARGE")

    def spill_evidence(self, *, queries_30=12000, epoch_start=None):
        """Spill-constrained golden fixture with the sample size / epoch age overridden (audit H2)."""
        payload = json.loads((self.golden / "spill-constrained.json").read_text())
        ev = payload["evidence"][0]
        epoch = next(x for x in ev["epochs"] if x["is_current"])
        epoch["query"]["queries_30"] = queries_30
        if epoch_start is not None:
            epoch["epoch_start"] = epoch_start
        return [p.evidence_from_dict(ev)]

    def size_candidate(self, evidence):
        w = p.evaluate_warehouses(evidence, self.t)[0]
        return w, next(c for c in w.candidates if c.knob == "warehouse_size")

    def test_upsize_requires_minimum_query_sample(self):
        # H2: 1 query that spills must not resize a warehouse.
        w, c = self.size_candidate(self.spill_evidence(queries_30=1))
        self.assertEqual(c.status, "INSUFFICIENT_EVIDENCE")
        self.assertIn("INSUFFICIENT_QUERY_SAMPLE", c.blockers)
        self.assertEqual(c.recommended_config, c.current_config)
        plan = next(s for s in w.scenarios if s.recommended)
        self.assertNotIn("warehouse_size", [s.knob for s in plan.steps])

    def test_upsize_requires_mature_config_epoch(self):
        # Analysis time is 2026-10-01T12:00Z, so this epoch is 2 days old (< 7).
        _, c = self.size_candidate(self.spill_evidence(epoch_start="2026-09-29T12:00:00+00:00"))
        self.assertEqual(c.status, "INSUFFICIENT_EVIDENCE")
        self.assertIn("CURRENT_CONFIG_EPOCH_TOO_NEW", c.blockers)
        self.assertEqual(c.recommended_config, c.current_config)

    def test_upsize_reports_every_missing_evidence_requirement(self):
        _, c = self.size_candidate(self.spill_evidence(queries_30=1, epoch_start="2026-09-29T12:00:00+00:00"))
        self.assertEqual(c.status, "INSUFFICIENT_EVIDENCE")
        self.assertEqual(set(c.blockers), {"INSUFFICIENT_QUERY_SAMPLE", "CURRENT_CONFIG_EPOCH_TOO_NEW"})

    def test_upsize_actionable_exactly_at_evidence_boundaries(self):
        # 100 queries and exactly 7 days of epoch age are sufficient.
        _, c = self.size_candidate(self.spill_evidence(queries_30=100, epoch_start="2026-09-24T12:00:00+00:00"))
        self.assertEqual(c.status, "ACTIONABLE")
        self.assertEqual(c.recommended_config["warehouse_size"], "X_LARGE")
        _, below = self.size_candidate(self.spill_evidence(queries_30=99))
        self.assertEqual(below.status, "INSUFFICIENT_EVIDENCE")

    def test_insufficient_evidence_upsize_has_no_savings_and_yields_valid_package(self):
        ev = self.spill_evidence(queries_30=1)
        _, c = self.size_candidate(ev)
        self.assertIsNone(c.estimated_savings_365)
        packages = p.build_recommendation_packages(ev, self.t)  # runs semantic validation
        pkg = packages[0]
        pkg_dict = pkg if isinstance(pkg, dict) else json.loads(p.canonical_json(pkg))
        size = next(x for x in pkg_dict["candidates"] if x["knob"] == "warehouse_size")
        self.assertEqual(size["status"], "INSUFFICIENT_EVIDENCE")
        self.assertEqual(size["independent_economics"]["estimation_status"], "UNAVAILABLE")
        self.assertEqual(size["recommended_value"], size["current_value"])

    # --- audit H3: severity is measured among *affected* queries, not all queries -------------------------
    NO_SEVERITY = {
        "spill_queries_30": 0, "p95_spill_bytes_when_spilled_30": None,
        "queued_queries_30": 0, "p95_wait_capacity_ms_when_queued_30": None,
        "compute_wait_queries_30": 0, "p95_wait_compute_ms_when_waited_30": None,
    }

    def severity_evidence(self, golden, **query):
        """Golden fixture re-expressed in the conditional-severity schema, with `query` overrides."""
        payload = json.loads((self.golden / f"{golden}.json").read_text())
        ev = payload["evidence"][0]
        epoch = next(x for x in ev["epochs"] if x["is_current"])
        for legacy in ("p95_spill_bytes_30", "p95_wait_capacity_ms_30", "p95_wait_compute_ms_30"):
            epoch["query"].pop(legacy, None)
        epoch["query"].update(self.NO_SEVERITY)
        epoch["query"].update(query)
        return [p.evidence_from_dict(ev)]

    def candidate(self, evidence, knob):
        return next(c for c in p.evaluate_warehouses(evidence, self.t)[0].candidates if c.knob == knob)

    GIB10 = 10 * 1024**3

    def test_rare_but_severe_spill_triggers_upsize(self):
        # 3% of 12,000 queries spill 10 GiB; the other 97% spill nothing. The old all-query p95 was 0.
        ev = self.severity_evidence("spill-constrained", spill_query_pct_30=3.0, spill_queries_30=360,
                                    p95_spill_bytes_when_spilled_30=self.GIB10)
        c = self.candidate(ev, "warehouse_size")
        self.assertEqual(c.status, "ACTIONABLE")
        self.assertEqual(c.recommended_config["warehouse_size"], "X_LARGE")

    def test_spill_severity_is_judged_among_spilled_queries(self):
        ev = self.severity_evidence("spill-constrained", spill_query_pct_30=3.0, spill_queries_30=360,
                                    p95_spill_bytes_when_spilled_30=200 * 1024**2)
        c = self.candidate(ev, "warehouse_size")
        self.assertNotEqual(c.status, "ACTIONABLE")
        self.assertIn("SPILL_PRESENT", c.blockers)

    def test_upsize_needs_minimum_affected_queries(self):
        # 100 queries is enough sample (H2), but 5 spilling queries cannot support a p95 severity claim.
        ev = self.severity_evidence("spill-constrained", queries_30=100, spill_query_pct_30=5.0, spill_queries_30=5,
                                    p95_spill_bytes_when_spilled_30=self.GIB10)
        c = self.candidate(ev, "warehouse_size")
        self.assertEqual(c.status, "INSUFFICIENT_EVIDENCE")
        self.assertIn("INSUFFICIENT_AFFECTED_QUERY_SAMPLE", c.blockers)
        self.assertEqual(c.recommended_config, c.current_config)

    def test_queue_severity_among_queued_queries_at_incidence_boundary(self):
        # Exactly 5.0% queued with a 45s tail among queued queries: increase signal.
        ev = self.severity_evidence("concurrency-constrained", queued_query_pct_30=5.0, queued_queries_30=600,
                                    p95_wait_capacity_ms_when_queued_30=45000)
        c = self.candidate(ev, "max_clusters")
        self.assertEqual(c.status, "ACTIONABLE")
        self.assertEqual(c.recommended_config["max_clusters"], 3)
        short = self.severity_evidence("concurrency-constrained", queued_query_pct_30=5.0, queued_queries_30=600,
                                       p95_wait_capacity_ms_when_queued_30=10000)
        self.assertNotEqual(self.candidate(short, "max_clusters").status, "ACTIONABLE")

    def test_max_clusters_increase_needs_minimum_affected_queries(self):
        ev = self.severity_evidence("concurrency-constrained", queries_30=100, queued_query_pct_30=10.0, queued_queries_30=10,
                                    p95_wait_capacity_ms_when_queued_30=45000)
        c = self.candidate(ev, "max_clusters")
        self.assertEqual(c.status, "INSUFFICIENT_EVIDENCE")
        self.assertIn("INSUFFICIENT_AFFECTED_QUERY_SAMPLE", c.blockers)

    def test_cold_start_waits_on_first_queries_block_auto_stop(self):
        # 30 first-after-start queries wait 120s out of 12,000; the all-query p95 compute wait was 0.
        ev = self.severity_evidence("cold-start-sensitive", compute_wait_queries_30=30,
                                    p95_wait_compute_ms_when_waited_30=120000)
        c = self.candidate(ev, "auto_stop_minutes")
        self.assertEqual(c.status, "BLOCKED")
        self.assertIn("COLD_START_SENSITIVE", c.blockers)

    def scaled(self, golden, *, max_clusters=None, **events):
        """Severity-schema golden fixture with the current epoch's max_clusters / events overridden (audit M3)."""
        import dataclasses
        ev = self.severity_evidence(golden)
        ep = ev[0].current_epoch
        if max_clusters is not None:
            ep.config = dataclasses.replace(ep.config, max_clusters=max_clusters)
        ep.events.update(events)
        return ev

    def queued(self, **overrides):
        return self.scaled("concurrency-constrained", **overrides)

    def increase_candidate(self, **events):
        ev = self.queued(**events)
        q = ev[0].current_epoch.query
        q.update(queued_query_pct_30=5.0, queued_queries_30=600, p95_wait_capacity_ms_when_queued_30=45000)
        return self.candidate(ev, "max_clusters")

    def test_cluster_recycling_blip_above_max_is_not_saturation(self):
        # Peak 3 > configured max 2 (recycle transient) but only 0.2% of running time actually at the ceiling.
        c = self.increase_candidate(max_observed_clusters=3, time_at_max_pct=0.2, minutes_at_max=20)
        self.assertEqual(c.status, "NO_CHANGE")
        self.assertEqual(c.recommended_config, c.current_config)

    def test_increase_needs_sustained_time_at_max_with_boundary(self):
        limit = self.t.time_at_max_pct_saturated
        self.assertEqual(limit, 5.0)
        up = self.increase_candidate(max_observed_clusters=2, time_at_max_pct=limit)
        self.assertEqual(up.status, "ACTIONABLE")
        self.assertEqual(up.recommended_config["max_clusters"], 3)
        below = self.increase_candidate(max_observed_clusters=2, time_at_max_pct=limit - 0.01)
        self.assertEqual(below.status, "NO_CHANGE")

    def reduction(self, max_clusters, peak):
        ev = self.scaled("idle-expensive-warehouse", max_clusters=max_clusters, max_observed_clusters=peak, time_at_max_pct=0.0, minutes_at_max=0.0)
        return self.candidate(ev, "max_clusters")

    def test_max_reduction_keeps_headroom_above_observed_peak(self):
        # Before: max 4, peak 3 -> target 3 (no growth room at all). Now one cluster of headroom is required.
        self.assertEqual(self.reduction(4, 3).status, "NO_CHANGE")
        self.assertEqual(self.reduction(4, 2).recommended_config["max_clusters"], 3)
        # Headroom scales with the peak: max(1, ceil(25% x peak)). Peak 8 needs 2 -> 10, which is not below max 10.
        self.assertEqual(self.reduction(10, 8).status, "NO_CHANGE")
        self.assertEqual(self.reduction(10, 4).recommended_config["max_clusters"], 5)
        self.assertEqual(self.reduction(10, 5).recommended_config["max_clusters"], 7)
        self.assertEqual(self.reduction(10, 2).recommended_config["max_clusters"], 3)

    def test_max_reduction_never_goes_below_min_clusters(self):
        c = self.reduction(10, 1)  # idle-expensive min_clusters is 2
        self.assertEqual(c.recommended_config["max_clusters"], 2)

    def test_few_long_capacity_waits_block_min_cluster_reduction(self):
        # Guard rules fire on any affected queries: 12 queued queries (0.1%) waiting 20s block a savings change.
        ev = self.severity_evidence("idle-expensive-warehouse", queued_query_pct_30=0.1, queued_queries_30=12,
                                    p95_wait_capacity_ms_when_queued_30=20000)
        c = self.candidate(ev, "min_clusters")
        self.assertEqual(c.status, "NO_CHANGE")
        self.assertIn("QUEUE_PRESSURE_PRESENT", c.blockers)

    def without_severity(self, golden, **query):
        """Golden fixture whose conditional-severity fields are absent (older collector), with `query` overrides."""
        payload = json.loads((self.golden / f"{golden}.json").read_text())
        ev = payload["evidence"][0]
        epoch = next(x for x in ev["epochs"] if x["is_current"])
        for key in list(self.NO_SEVERITY) + ["p95_spill_bytes_30", "p95_wait_capacity_ms_30", "p95_wait_compute_ms_30"]:
            epoch["query"].pop(key, None)
        epoch["query"].update(query)
        return [p.evidence_from_dict(ev)]

    def test_upsize_with_unknown_spill_severity_is_insufficient_evidence_not_a_silent_no_change(self):
        # Review finding: spill frequency crosses the threshold but severity is unknown. The decision must say so.
        for ev in (
            self.without_severity("spill-constrained"),  # count and p95 absent
            self.without_severity("spill-constrained", spill_queries_30=360, p95_spill_bytes_when_spilled_30=None),  # p95 null
        ):
            c = self.candidate(ev, "warehouse_size")
            self.assertEqual(c.status, "INSUFFICIENT_EVIDENCE")
            self.assertEqual(c.blockers, ["SPILL_SEVERITY_EVIDENCE_UNAVAILABLE"])
            self.assertEqual(c.recommended_config, c.current_config)

    def test_max_clusters_increase_with_unknown_wait_severity_is_insufficient_evidence(self):
        c = self.candidate(self.without_severity("concurrency-constrained"), "max_clusters")
        self.assertEqual(c.status, "INSUFFICIENT_EVIDENCE")
        self.assertEqual(c.blockers, ["QUEUE_WAIT_EVIDENCE_UNAVAILABLE"])
        self.assertEqual(c.recommended_config, c.current_config)

    def test_missing_severity_evidence_fails_closed(self):
        # Older evidence without the conditional-severity fields must not be read as 'no pressure' (zero).
        payload = json.loads((self.golden / "idle-expensive-warehouse.json").read_text())
        ev = payload["evidence"][0]
        epoch = next(x for x in ev["epochs"] if x["is_current"])
        for key in list(epoch["query"]):
            if key in self.NO_SEVERITY or key in ("p95_spill_bytes_30", "p95_wait_capacity_ms_30", "p95_wait_compute_ms_30"):
                del epoch["query"][key]
        evidence = [p.evidence_from_dict(ev)]
        minc = self.candidate(evidence, "min_clusters")
        self.assertEqual(minc.status, "NO_CHANGE")
        self.assertIn("QUEUE_WAIT_EVIDENCE_UNAVAILABLE", minc.blockers)
        auto = self.candidate(evidence, "auto_stop_minutes")
        self.assertEqual(auto.status, "BLOCKED")
        self.assertIn("COLD_START_EVIDENCE_UNAVAILABLE", auto.blockers)

    # --- H5: a withheld estimate is null, never zero --------------------------------------------------------
    def golden_package(self, name):
        payload = json.loads((self.golden / f"{name}.json").read_text())
        evidence = [p.evidence_from_dict(x) for x in payload["evidence"]]
        return p.build_recommendation_packages(evidence, self.t)[0].to_dict()

    @staticmethod
    def all_null(ranges):
        return all(ranges[k] is None for k in ("dbu", "aws", "total"))

    def test_blocked_candidate_economics_are_null_not_zero(self):
        pkg = self.golden_package("cold-start-sensitive")
        blocked = next(c for c in pkg["candidates"] if c["knob"] == "auto_stop_minutes")
        self.assertEqual(blocked["status"], "BLOCKED")
        econ = blocked["independent_economics"]
        self.assertEqual(econ["estimation_status"], "UNAVAILABLE")
        self.assertTrue(self.all_null(econ["savings_365d"]))
        self.assertTrue(self.all_null(econ["estimated_cost_365d"]))

    def test_insufficient_evidence_candidate_economics_are_null_not_zero(self):
        pkg = p.build_recommendation_packages(self.spill_evidence(queries_30=1), self.t)[0].to_dict()
        size = next(c for c in pkg["candidates"] if c["knob"] == "warehouse_size")
        self.assertEqual(size["status"], "INSUFFICIENT_EVIDENCE")
        self.assertEqual(size["independent_economics"]["estimation_status"], "UNAVAILABLE")
        self.assertTrue(self.all_null(size["independent_economics"]["savings_365d"]))

    def test_no_change_candidate_keeps_a_known_zero(self):
        pkg = self.golden_package("cold-start-sensitive")
        channel = next(c for c in pkg["candidates"] if c["knob"] == "warehouse_channel")
        econ = channel["independent_economics"]
        self.assertEqual(econ["estimation_status"], "NO_DELTA_MODELED")
        self.assertEqual(econ["savings_365d"]["total"], {"low": 0.0, "base": 0.0, "high": 0.0})

    def test_halted_scenario_economics_are_null_and_excluded_from_the_package(self):
        evidence = p.demo_evidence()
        evidence[0].sql_sku_prices = {}  # no target SQL price: the migration has no deterministic DBU counterfactual
        pkg = p.build_recommendation_packages(evidence, self.t)[0].to_dict()
        plan_b = next(s for s in pkg["scenarios"] if s["scenario_name"].startswith("Plan B"))
        econ = plan_b["portfolio_economics"]
        self.assertEqual(econ["estimation_status"], "UNAVAILABLE")
        for key in ("savings_365d", "estimated_cost_365d"):
            self.assertTrue(self.all_null(econ[key]), key)
        self.assertIsNone(econ["contingent_savings_365d"])
        self.assertTrue(all(st["projected_economics"] is None for st in plan_b["execution_plan"]))
        # Excluded from the envelope; the priced recommended plan still carries the package economics.
        self.assertIn(plan_b["scenario_id"], [x["scenario_id"] for x in pkg["package_economics"]["excluded_scenarios"]])
        self.assertIsNotNone(pkg["package_economics"]["estimated_savings_365d"])

    def test_scenario_with_no_steps_is_a_known_zero_not_unavailable(self):
        pkg = self.golden_package("newly-changed-config")
        plan_a = next(s for s in pkg["scenarios"] if s["scenario_name"].startswith("Plan A"))
        self.assertEqual([st for st in plan_a["execution_plan"] if st["step_type"] == "CANDIDATE"], [])
        econ = plan_a["portfolio_economics"]
        self.assertEqual(econ["estimation_status"], "NO_DELTA_MODELED")
        self.assertEqual(econ["savings_365d"]["total"], {"low": 0.0, "base": 0.0, "high": 0.0})

    def test_concurrency_constrained(self):
        _,r=self.load("concurrency-constrained"); c=next(c for c in r[0].candidates if c.knob=="max_clusters")
        self.assertEqual(c.status,"ACTIONABLE"); self.assertEqual(c.recommended_config["max_clusters"],3)

    def test_cold_start_sensitive(self):
        _,r=self.load("cold-start-sensitive"); c=next(c for c in r[0].candidates if c.knob=="auto_stop_minutes")
        self.assertEqual(c.status,"BLOCKED"); self.assertIn("COLD_START_SENSITIVE",c.blockers)

    def test_new_epoch_and_sparse_blockers(self):
        for name,blocker in (("newly-changed-config","CURRENT_CONFIG_EPOCH_TOO_NEW"),("sparse-evidence","INSUFFICIENT_QUERY_SAMPLE")):
            _,r=self.load(name); self.assertTrue(any(blocker in c.blockers for c in r[0].candidates),name)

    def test_photon_is_observed_and_reported_but_never_recommended(self):
        # Databricks publishes no on/off price for Photon on SQL warehouses, so the product leaves Photon at its
        # default: the billing-derived observation stays in current_config, but no candidate or plan step exists.
        for f in sorted(self.golden.glob("*.json")):
            payload=json.loads(f.read_text()); ev=[p.evidence_from_dict(x) for x in payload["evidence"]]
            for w in p.evaluate_warehouses(ev,self.t):
                self.assertNotIn("enable_photon",{c.knob for c in w.candidates},f.name)
                for sc in w.scenarios:
                    self.assertNotIn("enable_photon",[st.knob for st in sc.steps],f.name)
        pkg=self.golden_package("photon-disabled")
        self.assertEqual(pkg["current_config"]["photon"]["status"],"DISABLED")
        self.assertNotIn("enable_photon",[c["knob"] for c in pkg["candidates"]])
        self.assertEqual(len(pkg["candidates"]),6)
        plan_a=next(x for x in pkg["scenarios"] if x["scenario_name"].startswith("Plan A"))
        self.assertNotEqual(plan_a["portfolio_economics"]["estimation_status"],"UNAVAILABLE")  # no longer blocked by Photon
        self.assertIsNotNone(plan_a["portfolio_economics"]["savings_365d"]["total"])

    def test_serverless_to_pro_costed_with_workspace_calibration(self):
        _,r=self.load("serverless-to-pro-costed"); w=self.warehouse(r,"wh-serverless-001")
        c=next(c for c in w.candidates if c.knob=="warehouse_type"); self.assertEqual(c.recommended_config,{"warehouse_type":"PRO"})
        self.assertIsNotNone(c.dbu_365d_estimated_cost_range); self.assertIsNotNone(c.aws_365d_estimated_cost_range)
        self.assertEqual(c.metadata["migration_model"]["dbu_rate_source"],"WORKSPACE_SYSTEM_TABLE_EMPIRICAL_PRO_CLASSIC_DBU_PER_CLUSTER_HOUR")
        b=next(s for s in w.scenarios if s.name.startswith("Plan B")); self.assertEqual(b.portfolio_estimate.estimate_status,"PRE_CANARY_MODELED_RANGE")

    def test_missing_aws_region_is_explicit(self):
        _,r=self.load("aws-region-missing"); w=r[0]
        self.assertIsNone(w.aws_365d_baseline_cost); self.assertIsNone(w.total_365d_baseline_cost)



if __name__ == "__main__":
    unittest.main()
