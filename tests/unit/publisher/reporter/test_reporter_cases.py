import json
import tempfile
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))
from publisher.config.versions import PUBLISHER_VERSION  # noqa: E402
from publisher.reporter.json_reporter import package_locked_contract  # noqa: E402
from tests.publisher_case import PublisherCase  # noqa: E402


class ReporterCases(PublisherCase):
    def test_reporter_json_workbook_and_hitl_findings(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            packages = self.packages
            files = package_locked_contract(Path(td), packages, self.evidence)
            payload = json.loads(Path(files[f"package:{self.evidence[0].workspace_id}:{self.evidence[0].warehouse_id}"]).read_text())
            for field in ("candidates", "scenarios", "recommended_plan", "package_economics", "evidence", "workload_profile", "realization"):
                self.assertIn(field, payload)
            self.assertEqual(len(payload["candidates"]), 5)
            self.assertGreater(len(payload["evidence"]), 0)
            wb = load_workbook(files["workbook"])
            self.assertEqual(wb.sheetnames, ["Executive Summary", "Warehouse Portfolio", "Scenario Plans", "Recommendation Registry"])

    def test_candidate_and_recommended_plan_views(self):
        from publisher.contracts import workbook_contract as wc
        projection = wc.project_workbook(self.packages)
        row = projection.warehouse_portfolio[0]
        self.assertNotIn("Photon", row["candidate_view"])
        self.assertIn("Size", row["candidate_view"])
        self.assertIn("Type", row["candidate_view"])
        self.assertIn("warehouse_size", row["recommended_execution_plan"])
        self.assertIn("READY_FOR_CANARY", row["recommended_execution_plan"])
        self.assertIn("remeasure 7d", row["recommended_execution_plan"])

    def test_manifest_carries_locked_contract_and_workbook_metadata(self):
        with tempfile.TemporaryDirectory() as td:
            packages = self.packages
            files = package_locked_contract(Path(td), packages, self.evidence)
            manifest = json.loads(Path(files["manifest"]).read_text())
            self.assertEqual(manifest["publisher_version"], PUBLISHER_VERSION)
            self.assertEqual(manifest["reporter_contract_version"], "3.1.0")
            self.assertEqual(manifest["workbook_contract_version"], "1.6.0")
            self.assertEqual(manifest["workbook_status"], "GENERATED")
            self.assertEqual(manifest["package_count"], 1)
            self.assertEqual(manifest["workbook_file"], "workbook.xlsx")
