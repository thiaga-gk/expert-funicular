import json
import tempfile
import unittest
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT))
sys.path.insert(0,str(ROOT / "src"))
from publisher import publisher as p
from publisher.analyzer.collector.demo import DEMO_FLEET, demo_evidence, synthetic_demo_evidence  # noqa: E402
from publisher.governor.thresholds import Thresholds  # noqa: E402
from publisher.reporter.json_reporter import package_locked_contract  # noqa: E402
from publisher.contracts import reporter_contract as rc  # noqa: E402
from publisher.contracts import workbook_contract as wc
from publisher.reporter import workbook_reporter as wr


class WorkbookContractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.t=Thresholds()
        cls.evidence=synthetic_demo_evidence()
        cls.packages=p.build_recommendation_packages(cls.evidence,cls.t)
        cls.projection=wc.project_workbook(cls.packages)

    def test_four_locked_sheet_grains(self):
        self.assertEqual(wc.SHEET_NAMES,("Executive Summary","Warehouse Portfolio","Scenario Plans","Recommendation Registry"))
        n=len(DEMO_FLEET)
        self.assertEqual(len(self.packages),n)
        self.assertEqual(len(self.projection.warehouse_portfolio),n)
        self.assertEqual(len(self.projection.scenario_plans),2*n)
        self.assertEqual(len(self.projection.recommendation_registry),5*n)

    def test_scenario_plans_carry_contingent_upside_apart_from_authoritative_impact(self):
        """Contract 1.1.0 / workbook 1.2.0: contingent upside is separate, numeric, signed and tagged in the plan text."""
        rows = self.projection.scenario_plans
        keys = ("contingent_total_low_impact", "contingent_total_base_impact", "contingent_total_high_impact")
        self.assertTrue(set(keys) <= {c.key for c in wc.SCENARIO_COLUMNS})
        priced = [r for r in rows if isinstance(r["contingent_total_base_impact"], (int, float))]
        self.assertTrue(priced)  # e.g. a size canary with min/max/auto-stop pending
        for r in rows:
            if r in priced:
                self.assertIn("PENDING_REEVALUATION", r["execution_plan"])
                self.assertIn("priced CONTINGENT", r["execution_plan"])
                for k in keys:
                    self.assertIsInstance(r[k], (int, float))
                    self.assertNotIsInstance(r[k], bool)
            else:  # nothing pending, or pending but unpriceable: unavailable, never zero
                for k in keys:
                    self.assertEqual(r[k], wc.DISPLAY_UNAVAILABLE)

    def test_workbook_rows_trace_back_to_locked_ids(self):
        self.assertEqual({x["recommendation_package_id"] for x in self.projection.warehouse_portfolio},{p.recommendation_package_id for p in self.packages})
        self.assertTrue(all(x["recommendation_package_id"] for x in self.projection.scenario_plans))
        self.assertTrue(all(x["recommendation_package_id"] and x["candidate_id"] and x["technique_id"] for x in self.projection.recommendation_registry))

    def test_candidate_registry_has_hitl_evidence_and_nonadditive_economics(self):
        self.assertTrue(all(x["evidence_json"] and x["evidence_refs_json"] for x in self.projection.recommendation_registry))
        self.assertTrue(all(x["is_additive"] is False for x in self.projection.recommendation_registry))

    def test_portfolio_matches_exact_locked_columns_and_package_economics(self):
        row=self.projection.warehouse_portfolio[0]
        self.assertEqual(set(row),{c.key for c in wc.PORTFOLIO_COLUMNS})
        for key in ("candidate_view","recommended_execution_plan","package_low_impact","package_base_impact","package_high_impact","package_cost_low_365d","package_cost_base_365d","package_cost_high_365d","active_days_30d","query_source_profile_json","scenario_ids_json","candidate_ids_json"):
            self.assertIn(key,row)
        self.assertNotIn("Photon",row["candidate_view"])

    def test_scenario_and_registry_match_exact_locked_columns(self):
        self.assertEqual(set(self.projection.scenario_plans[0]),{c.key for c in wc.SCENARIO_COLUMNS})
        self.assertEqual(set(self.projection.recommendation_registry[0]),{c.key for c in wc.REGISTRY_COLUMNS})
        for key in ("dbu_low_impact","dbu_base_impact","dbu_high_impact","aws_low_impact","aws_base_impact","aws_high_impact","total_low_impact","total_base_impact","total_high_impact","estimated_total_cost_low","estimated_total_cost_base","estimated_total_cost_high"):
            self.assertIn(key,self.projection.scenario_plans[0])
            self.assertIn(key,self.projection.recommendation_registry[0])

    def test_presentation_has_no_blank_semantic_cells_and_impacts_are_numeric(self):
        collections=(
            self.projection.executive_metrics, self.projection.executive_opportunities,
            self.projection.warehouse_portfolio, self.projection.scenario_plans, self.projection.recommendation_registry,
        )
        for rows in collections:
            for row in rows:
                self.assertTrue(all(v is not None and (not isinstance(v,str) or v != "") for v in row.values()))
        impact_keys = [c.key for c in wc.PORTFOLIO_COLUMNS + wc.SCENARIO_COLUMNS + wc.REGISTRY_COLUMNS if "impact" in c.key and "type" not in c.key and c.number_format]
        for rows in (self.projection.warehouse_portfolio,self.projection.scenario_plans,self.projection.recommendation_registry):
            for row in rows:
                for key in impact_keys:
                    if key in row and row[key] != wc.DISPLAY_UNAVAILABLE:
                        self.assertIsInstance(row[key], (int,float), key)
                        self.assertNotIsInstance(row[key], bool, key)
        for rows in collections:
            for row in rows:
                for value in row.values():
                    if isinstance(value,str):
                        self.assertFalse(value.startswith("Save USD"))
                        self.assertFalse(value.startswith("Cost +USD"))

    def test_identifier_columns_are_text_formatted(self):
        for columns in (wc.PORTFOLIO_COLUMNS,wc.SCENARIO_COLUMNS,wc.REGISTRY_COLUMNS):
            for spec in columns:
                if spec.key in {"recommendation_package_id","workspace_id","warehouse_id","current_epoch_id","recommended_scenario_id","scenario_id","candidate_id","technique_id","epoch_id"}:
                    self.assertEqual(spec.number_format,"@",spec.key)

    def test_workbook_contract_column_counts_and_long_identifier_rendering(self):
        self.assertEqual((len(wc.PORTFOLIO_COLUMNS),len(wc.SCENARIO_COLUMNS),len(wc.REGISTRY_COLUMNS)),(61,45,56))
        # Long digit-only identifiers are written as explicit text cells, so no apostrophe workaround is needed.
        self.assertEqual(wr._excel_text("1234567890123456"), "1234567890123456")
        self.assertEqual(wr._excel_text("wh-synth-001"), "wh-synth-001")

    def test_rendered_workbook_satisfies_the_contract(self):
        """Render the real .xlsx and validate sheets, columns, grains, text IDs, numeric economics, no formulas."""
        from openpyxl import load_workbook
        from scripts.validate_workbook import validate_workbook
        with tempfile.TemporaryDirectory() as td:
            path = wr.write_workbook(Path(td) / "workbook.xlsx", self.packages)
            packages = [json.loads(json.dumps(pkg.to_dict())) for pkg in self.packages]
            self.assertEqual(validate_workbook(path, packages), [])
            wb = load_workbook(path)
            portfolio = wb[wc.SHEET_PORTFOLIO]
            headers = [c.value for c in portfolio[4]]
            ws_col = headers.index("Workspace ID") + 1
            cell = portfolio.cell(row=5, column=ws_col)
            self.assertEqual(cell.data_type, "s")          # a long digit-only id stays exact text
            self.assertEqual(cell.value, self.packages[0].warehouse.workspace_id)
            scenarios = wb[wc.SHEET_SCENARIOS]
            sh = [c.value for c in scenarios[4]]
            base = sh.index("Total base impact") + 1
            self.assertTrue(any(isinstance(scenarios.cell(r, base).value, float) for r in range(5, scenarios.max_row + 1)))

    def test_workbook_validator_rejects_formulas_string_money_and_wrong_columns(self):
        from openpyxl import load_workbook
        from scripts.validate_workbook import validate_workbook
        with tempfile.TemporaryDirectory() as td:
            path = wr.write_workbook(Path(td) / "workbook.xlsx", self.packages[:2])
            self.assertEqual(validate_workbook(path), [])
            wb = load_workbook(path)
            scenarios = wb[wc.SHEET_SCENARIOS]
            base = [c.value for c in scenarios[4]].index("Total base impact") + 1
            scenarios.cell(row=5, column=base).value = "Save USD 1,000"          # string money
            wb[wc.SHEET_PORTFOLIO]["A5"].value = "=SUM(1,2)"                      # business formula
            wb[wc.SHEET_REGISTRY]["B4"].value = "Renamed column"                   # contract column drift
            wb[wc.SHEET_EXECUTIVE].sheet_state = "hidden"                          # hidden sheet
            tampered = Path(td) / "tampered.xlsx"
            wb.save(tampered)
            joined = "\n".join(validate_workbook(tampered))
            for expected in ("must be numeric", "formulas are not allowed", "header row differs", "must be visible"):
                self.assertIn(expected, joined)

    def test_text_that_looks_like_a_formula_is_written_as_text(self):
        from openpyxl import Workbook
        sheet = Workbook().active
        cell = wr._put(sheet, 1, 1, "=1+1")
        self.assertEqual(cell.data_type, "s")

    def test_oversized_audit_cells_are_truncated_with_a_visible_marker(self):
        from openpyxl import Workbook
        sheet = Workbook().active
        cell = wr._put(sheet, 1, 1, "x" * 40000)
        self.assertLessEqual(len(cell.value), wr.EXCEL_MAX_CELL_CHARS)
        self.assertTrue(cell.value.endswith(wr.TRUNCATION_MARKER))

    def test_package_writer_places_workbook_in_output_dir(self):
        with tempfile.TemporaryDirectory() as td:
            files=package_locked_contract(Path(td),self.packages[:2],self.evidence[:2])
            self.assertEqual(Path(files["workbook"]).name,"workbook.xlsx")
            self.assertTrue(Path(files["workbook"]).exists())
            manifest=json.loads(Path(files["manifest"]).read_text())
            self.assertEqual(manifest["workbook_contract_version"],"1.4.0")
            self.assertEqual(manifest["workbook_status"],"GENERATED")




class PlanStepSavingsAndExecutiveRowsTests(unittest.TestCase):
    """Reporter 1.3.0 per-step savings and the Workbook 1.4.0 plan-vs-experiment and exclusion rows."""

    @classmethod
    def setUpClass(cls):
        cls.t=Thresholds()
        cls.packages=p.build_recommendation_packages(synthetic_demo_evidence(),cls.t)

    def test_step_savings_are_ordered_and_add_up_along_the_authoritative_chain(self):
        checked=0
        for pkg in self.packages:
            plan=next(s for s in pkg.scenarios if s.status==rc.ScenarioStatus.RECOMMENDED)
            auth=[x.projected_economics for x in plan.execution_plan if x.projected_economics is not None
                  and x.projected_economics.pricing_basis==rc.PricingBasis.AUTHORITATIVE and x.projected_economics.step_savings_365d is not None]
            for pe in auth:
                total=pe.step_savings_365d.total
                if total is not None:
                    self.assertTrue(total.low<=total.base<=total.high)
            if len(auth)>=2 and all(x.step_savings_365d.total is not None for x in auth):
                self.assertAlmostEqual(sum(x.step_savings_365d.total.base for x in auth), auth[-1].running_savings_365d.total.base, delta=0.05*len(auth))
                self.assertAlmostEqual(auth[0].step_savings_365d.total.base, auth[0].running_savings_365d.total.base, delta=0.03)
                checked+=1
        self.assertGreater(checked,0)

    def test_unpriced_step_has_no_step_savings(self):
        for pkg in self.packages:
            for s in pkg.scenarios:
                for x in s.execution_plan:
                    if x.projected_economics is None:
                        continue
                    if x.projected_economics.running_savings_365d.total is None:
                        self.assertTrue(x.projected_economics.step_savings_365d is None or x.projected_economics.step_savings_365d.total is None)

    def metrics(self, excluded):
        return {m["metric"]: m for m in wc.project_workbook(self.packages, excluded).executive_metrics}

    def test_executive_summary_separates_plan_savings_from_experiment_upside(self):
        m=self.metrics([])
        low,base,high=(m[f"Plan savings {b}"]["value"] for b in ("low","base","high"))
        self.assertTrue(low<=base<=high)
        self.assertGreater(m["Migration experiment high"]["value"], 0)
        self.assertEqual(m["Excluded (unsupported type)"]["value"], 0)

    def test_cost_notes_state_the_pricing_basis(self):
        m=self.metrics([])
        self.assertIn("DBU at list price", m["Forward 365d baseline"]["notes"])
        self.assertIn("on-demand EC2", m["Forward 365d baseline"]["notes"])
        for b in ("low","base","high"):
            self.assertIn("DBU at list price, AWS on-demand", m[f"Plan savings {b}"]["notes"])

    def test_package_assumptions_state_the_pricing_basis(self):
        text=" ".join(self.packages[0].cost_baseline.forecast.assumptions)
        self.assertIn("list price; negotiated discounts lower real savings in proportion", text)
        self.assertIn("on-demand EC2 rates; Spot, Reserved Instances and Savings Plans are not modeled", text)

    def test_exclusion_row_counts_unsupported_warehouses_or_says_unavailable(self):
        excluded=[{"workspace_id":"w","warehouse_id":"x","warehouse_type":"LAKEHOUSE_REALTIME","reason":"UNSUPPORTED_WAREHOUSE_TYPE"}]
        row=self.metrics(excluded)["Excluded (unsupported type)"]
        self.assertEqual(row["value"],1)
        self.assertIn("LAKEHOUSE_REALTIME",row["notes"])
        self.assertEqual(self.metrics(None)["Excluded (unsupported type)"]["value"], wc.DISPLAY_UNAVAILABLE)

    def test_plan_text_states_each_priced_steps_own_savings(self):
        texts=[r.get("execution_plan") or "" for r in wc.project_workbook(self.packages, []).scenario_plans]
        self.assertTrue(any("; step saves " in t for t in texts))

    def test_incompatible_migration_is_not_experiment_upside(self):
        e=demo_evidence()[0]
        e.compatibility_evidence={"mode":"B_READ_ONLY_API","status":"AVAILABLE","enable_serverless_compute":False,"external_legacy_hive_metastore_detected":True}
        pkgs=p.build_recommendation_packages([e],self.t)
        plan_b=next(s for s in pkgs[0].scenarios if s.scenario_type==rc.ScenarioType.WAREHOUSE_TYPE_MIGRATION)
        self.assertEqual(plan_b.compatibility.status,rc.CompatibilityStatus.INCOMPATIBLE)
        self.assertGreater(plan_b.portfolio_economics.savings_365d.total.high,0)
        row={m["metric"]:m for m in wc.project_workbook(pkgs,[]).executive_metrics}["Migration experiment high"]
        self.assertEqual(row["value"],0)


if __name__ == "__main__":
    unittest.main()
