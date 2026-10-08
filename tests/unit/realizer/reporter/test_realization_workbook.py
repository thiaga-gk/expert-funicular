"""The realization workbook (src/realizer/contracts/report.py, the projection; src/realizer/reporter/workbook.py, the renderer).

Realistic inputs: real Publisher packages realized into a temporary output directory, their final realization packages,
ledgers and state index read back and handed over as ``WarehouseRun``s, the way a fleet run hands them over.

* wh-demo-039: the balanced and aggressive cost-profile packages of one epoch (the end-to-end goldens ge07 and ge06) and one
  change (auto-stop 210 on 2026-10-08 14:32). The balanced lineage claims it; the aggressive one is observable only and
  names the balanced package.
* wh-demo-013 (GE-01's package, auto-stop 30 to 10): applied 2026-10-08 and rolled back to 30 on 2026-11-05; and, in
  another output directory, never applied.
* A refused and a crashed warehouse, with no realization.
"""
import copy
import dataclasses
import json
import tempfile
import unittest
from pathlib import Path
from typing import Dict, List, Mapping, Sequence, Tuple

from openpyxl import load_workbook

from tests.unit.realizer.helpers import canonical, fixture_package  # noqa: F401  (src/realizer, never a shadowing copy)
from tests import realizer_world as world
from tests.golden.realizer import e2e
from tests.unit.realizer.lineage_support import AGGRESSIVE, BALANCED, stored_run
from realizer.analyzer.projector import project_observation
from realizer.contracts import realization as rc
from realizer.contracts import report
from realizer.contracts.report import WarehouseRun
from realizer.packager.ledger import ledger_path, load_ledger
from realizer.realizer import realize
from realizer.reporter.workbook import write_realization_workbook

GE01 = "ge01-auto-stop-exact-scoped"
ROLLBACK = [{"at": "2026-10-08T00:00:00Z", "set": {"auto_stop": 10}}, {"at": "2026-11-05T09:00:00Z", "set": {"auto_stop": 30}}]


def warehouse_run(out: Path, package: Mapping, *, status: str = report.RUN_STATUS_OK, with_index: bool = True) -> WarehouseRun:
    """The ``WarehouseRun`` a fleet run hands over for *package*'s warehouse: every lineage of it in *out* (the final
    realization packages, in package id order), their ledgers and the state index."""
    wh = package["warehouse"]
    realizations = []
    for path in sorted((out / "realization" / "packages").glob("*.json")):
        body = json.loads(path.read_text(encoding="utf-8"))
        if body["warehouse"]["warehouse_id"] == wh["warehouse_id"]:
            realizations.append(rc.RealizationPackage(**body))
    ledgers = {r.realization_package_id: load_ledger(ledger_path(out, r.realization_package_id)) for r in realizations}
    index = None
    if with_index:
        (path,) = sorted((out / "realization" / "state").glob(f"*__{wh['warehouse_id']}.json"))
        index = json.loads(path.read_text(encoding="utf-8"))
    return WarehouseRun(workspace_id=wh["workspace_id"], warehouse_id=wh["warehouse_id"], status=status,
                        recommendation_package_id=package["recommendation_package_id"],
                        realizations=tuple(realizations), ledgers=ledgers, state_index=index)


def ge01_observation(package: Mapping, changes: Sequence[Mapping], day: str) -> Dict:
    spec = e2e.world_spec(13, changes)
    return project_observation(package, world.measure(package, spec, 13, as_of=f"{day}T12:00:00+00:00"))


def by_rpkg(rows: Sequence[Mapping], rpkg: str) -> Dict:
    (row,) = [r for r in rows if r["realization_package_id"] == rpkg]
    return dict(row)


class Fixture:
    """The realized inputs, built once for the module."""
    pair: WarehouseRun
    rolled_back: WarehouseRun
    not_applied: WarehouseRun
    balanced_pkg: str
    aggressive_pkg: str


def setUpModule():
    aggressive, balanced = stored_run(AGGRESSIVE), stored_run(BALANCED)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        for package, observation in (aggressive, balanced):
            realize(copy.deepcopy(package), copy.deepcopy(observation), out)
        Fixture.pair = warehouse_run(out, balanced[0])
    Fixture.balanced_pkg, Fixture.aggressive_pkg = balanced[0]["recommendation_package_id"], aggressive[0]["recommendation_package_id"]
    ge01 = Fixture.ge01 = e2e.stored_package(GE01)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        for day in ("2026-10-23", "2026-11-04", "2026-11-20"):
            realize(copy.deepcopy(ge01), ge01_observation(ge01, ROLLBACK, day), out)
        Fixture.rolled_back = warehouse_run(out, ge01)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        realize(copy.deepcopy(ge01), ge01_observation(ge01, [], "2026-10-07"), out)
        Fixture.not_applied = warehouse_run(out, ge01)


REFUSED = WarehouseRun(workspace_id="1234567890123456", warehouse_id="wh-demo-020", status=report.RUN_STATUS_REFUSED,
                       reason="package PKG-0123456789ABCDEF01234567 fails the schema", recommendation_package_id="PKG-0123456789ABCDEF01234567")
CRASHED = WarehouseRun(workspace_id="1234567890123456", warehouse_id="wh-demo-021", status=report.RUN_STATUS_CRASHED, reason="KeyError")


def keys(columns) -> List[str]:
    return [c.key for c in columns]


class ContractTests(unittest.TestCase):
    def test_version_and_sheets(self):
        self.assertEqual(report.REALIZATION_WORKBOOK_CONTRACT_VERSION, "2.0.0")
        self.assertEqual(report.SHEET_NAMES, ("Warehouses", "Lineages", "Timeline"))
        self.assertEqual(tuple(report.SHEET_COLUMNS), report.SHEET_NAMES)

    def test_columns_are_unique_grouped_and_the_audit_json_is_at_the_far_right(self):
        for sheet, (columns, groups) in report.SHEET_COLUMNS.items():
            with self.subTest(sheet=sheet):
                names = keys(columns)
                self.assertEqual(len(names), len(set(names)))
                self.assertEqual(len({c.header for c in columns}), len(columns))
                # The groups cover every column once, left to right.
                covered = []
                for group in groups:
                    covered.extend(names[names.index(group.start_key): names.index(group.end_key) + 1])
                self.assertEqual(covered, names)
                json_keys = [k for k in names if k.endswith("_json") or k == "evidence_refs"]
                self.assertTrue(json_keys)
                self.assertEqual(names[-len(json_keys):], json_keys)
                self.assertEqual(groups[-1].title, "Audit")

    def test_warehouse_columns(self):
        self.assertEqual(keys(report.WAREHOUSE_COLUMNS), [
            "workspace_id", "warehouse_id", "warehouse_name", "recommendation_package_id", "realization_package_id", "lineage_basis",
            "run_status", "run_reason",
            "applied_knob", "from_value", "applied_value", "recommended_value", "applied_at_utc", "applied_until_utc",
            "implementation_conformance", "validation_state", "realization_state", "maturity_state",
            "currency", "realized_to_date", "days_covered", "first_day", "last_day", "annualized_realized_savings",
            "expected_same_scope_annual_savings", "realization_ratio",
            "attribution_claim_status", "claimed_by_package_id", "other_lineages",
            "next_action", "as_of_utc",
            "application_json", "economics_json", "attribution_json", "applied_knob_owners_json",
        ])

    def test_lineage_columns_are_the_summary_columns_plus_the_realization_columns(self):
        summary = list(report.project_summary(Fixture.pair.realizations[0]))
        added = ["claimed_by_package_id", "currency", "realized_to_date", "realized_to_date_days_covered", "next_action"]
        self.assertEqual(sorted(keys(report.LINEAGE_COLUMNS)), sorted(summary + added))
        names = keys(report.LINEAGE_COLUMNS)
        self.assertEqual(names[-4:], ["application_json", "validation_json", "economics_json", "attribution_json"])

    def test_timeline_columns_are_the_ledger_projection(self):
        lineage = Fixture.pair.realizations[0]
        rows = report.project_timeline(Fixture.pair.ledgers[lineage.realization_package_id])
        self.assertEqual(keys(report.TIMELINE_COLUMNS), list(rows[0]))


class NextActionTests(unittest.TestCase):
    """Every realization state has a label; regeneration status and plan progress refine it where they matter."""

    def lineage(self, state: str, *, regeneration: str = "NOT_REQUIRED", actionable: int = 1) -> rc.RealizationPackage:
        base = Fixture.not_applied.realizations[0]
        application = {**base.application, "plan_progress": {**base.application["plan_progress"], "authoritative_actionable_steps": actionable}}
        return dataclasses.replace(base, current_state={**base.current_state, "realization_state": state},
                                   regeneration={**base.regeneration, "status": regeneration}, application=application)

    def test_an_applied_change_validated_without_measured_savings_keeps_monitoring(self):
        # A size change: the guards passed, but no counterfactual measures its savings yet (economics unavailable).
        r = self.lineage("APPLIED")
        r = dataclasses.replace(r, validation={**r.validation, "state": "PASSED"}, economics={**r.economics, "available": False})
        self.assertEqual(report.next_action(r), "Keep monitoring: savings not measured for this change yet")
        waiting = dataclasses.replace(r, validation={**r.validation, "state": "NOT_READY"})
        self.assertEqual(report.next_action(waiting), "Wait for validation")

    def test_every_realization_state(self):
        expected = {
            "NOT_APPLIED": "Apply the recommended change",
            "APPLIED": "Wait for validation",
            "VALIDATING": "Wait for validation",
            "REALIZED": "Keep monitoring",
            "MONITORING": "Keep monitoring",
            "PERFORMANCE_VALIDATED": "Keep monitoring",
            "CAPACITY_VALIDATED": "Keep monitoring",
            "SAFETY_ACTION_APPLIED": "Keep monitoring: safety rollback in force",
            "UNDER_REALIZED": "Review: re-optimization recommended",
            "VALIDATION_FAILED": "Roll back: a guard failed",
            "ROLLED_BACK": "None: rolled back",
            "DRIFTED": "Run the Publisher: regeneration requested",
            "REGENERATING": "Run the Publisher: regeneration requested",
            "BLOCKED": "Review the blocker",
            "INVALIDATED": "Review the blocker",
        }
        self.assertEqual(set(expected), {s.value for s in rc.RealizationState})
        for state, label in expected.items():
            with self.subTest(state=state):
                self.assertEqual(report.next_action(self.lineage(state)), label)

    def test_regeneration_status_and_plan_progress(self):
        self.assertEqual(report.next_action(self.lineage("ROLLED_BACK", regeneration="REQUESTED")), "Run the Publisher: regeneration requested")
        for state in ("ROLLED_BACK", "DRIFTED", "REGENERATING"):
            with self.subTest(state=state):
                self.assertEqual(report.next_action(self.lineage(state, regeneration="REPLACED")), "None: replacement package linked")
        self.assertEqual(report.next_action(self.lineage("NOT_APPLIED", actionable=0)), "None: no step to apply")

    def test_the_table_holds_every_label_once_per_rule(self):
        labels = set(report.NEXT_ACTIONS.values())
        self.assertIn("Fix the run: REFUSED", labels)
        self.assertIn("Fix the run: CRASHED", labels)
        self.assertEqual({state for state, _ in report.NEXT_ACTIONS if state not in report.RUN_STATUSES and state != report.NO_LINEAGE},
                         {s.value for s in rc.RealizationState})

    def test_an_unknown_state_fails_closed(self):
        with self.assertRaises(ValueError):
            report.next_action(self.lineage("SOMETHING_NEW"))


class WarehouseRowTests(unittest.TestCase):
    def test_the_lineage_that_claims_the_change_is_the_warehouse_row(self):
        (row,) = report.project_realization_workbook([Fixture.pair]).warehouses
        balanced = by_rpkg(report.project_realization_workbook([Fixture.pair]).lineages, rc.make_rpkg_id(Fixture.balanced_pkg))
        self.assertEqual({k: row[k] for k in ("warehouse_id", "recommendation_package_id", "realization_package_id", "lineage_basis",
                                              "run_status", "run_reason", "applied_knob", "from_value", "applied_value", "recommended_value",
                                              "applied_at_utc", "applied_until_utc", "implementation_conformance", "validation_state",
                                              "realization_state", "maturity_state", "currency", "realized_to_date", "days_covered",
                                              "first_day", "last_day", "annualized_realized_savings", "expected_same_scope_annual_savings",
                                              "attribution_claim_status", "claimed_by_package_id", "other_lineages", "next_action", "as_of_utc")}, {
            "warehouse_id": "wh-demo-039", "recommendation_package_id": Fixture.balanced_pkg,
            "realization_package_id": rc.make_rpkg_id(Fixture.balanced_pkg), "lineage_basis": "CLAIMS_LATEST_CHANGE",
            "run_status": "OK", "run_reason": None, "applied_knob": "auto_stop_minutes", "from_value": 365, "applied_value": 210,
            "recommended_value": 210, "applied_at_utc": "2026-10-08T14:32:00Z", "applied_until_utc": None,
            "implementation_conformance": "EXACT", "validation_state": "PASSED", "realization_state": "REALIZED",
            "maturity_state": "REPRESENTATIVE", "currency": "USD", "realized_to_date": 1069.73, "days_covered": 28.0,
            "first_day": "2026-10-09", "last_day": "2026-11-05", "annualized_realized_savings": 13944.69,
            "expected_same_scope_annual_savings": balanced["expected_same_scope_annual_savings"],
            "attribution_claim_status": "CLAIMABLE", "claimed_by_package_id": None, "other_lineages": 1,
            "next_action": "Keep monitoring", "as_of_utc": "2026-11-07T12:00:00Z",
        })
        self.assertIsInstance(row["realization_ratio"], float)
        self.assertEqual(json.loads(row["applied_knob_owners_json"]), Fixture.pair.state_index["applied_knob_owners"])

    def test_without_a_state_index_the_claimable_lineage_with_the_latest_application_is_chosen(self):
        run = dataclasses.replace(Fixture.pair, state_index=None)
        (row,) = report.project_realization_workbook([run]).warehouses
        self.assertEqual((row["recommendation_package_id"], row["lineage_basis"], row["applied_knob_owners_json"]),
                         (Fixture.balanced_pkg, "CLAIMS_LATEST_CHANGE", None))

    def test_a_lineage_another_lineage_claims_names_the_owner(self):
        aggressive = [r for r in Fixture.pair.realizations if r.source_recommendation["recommendation_package_id"] == Fixture.aggressive_pkg]
        run = dataclasses.replace(Fixture.pair, realizations=tuple(aggressive), state_index=None)
        (row,) = report.project_realization_workbook([run]).warehouses
        self.assertEqual({k: row[k] for k in ("recommendation_package_id", "lineage_basis", "implementation_conformance",
                                              "attribution_claim_status", "claimed_by_package_id", "other_lineages", "realized_to_date")}, {
            "recommendation_package_id": Fixture.aggressive_pkg, "lineage_basis": "NEWEST_ACTIVE_PACKAGE",
            "implementation_conformance": "MODIFIED", "attribution_claim_status": "OBSERVABLE_ONLY",
            "claimed_by_package_id": Fixture.balanced_pkg, "other_lineages": 0, "realized_to_date": 1069.73})

    def test_a_rolled_back_lineage(self):
        (row,) = report.project_realization_workbook([Fixture.rolled_back]).warehouses
        self.assertEqual({k: row[k] for k in ("warehouse_id", "lineage_basis", "applied_knob", "from_value", "applied_value",
                                              "applied_at_utc", "applied_until_utc", "realization_state", "realized_to_date",
                                              "days_covered", "first_day", "last_day", "annualized_realized_savings",
                                              "expected_same_scope_annual_savings", "realization_ratio", "attribution_claim_status",
                                              "next_action", "as_of_utc")}, {
            "warehouse_id": "wh-demo-013", "lineage_basis": "CLAIMS_LATEST_CHANGE", "applied_knob": "auto_stop_minutes",
            "from_value": 30, "applied_value": 10, "applied_at_utc": "2026-10-08T00:00:00Z", "applied_until_utc": "2026-11-05T09:00:00Z",
            "realization_state": "ROLLED_BACK", "realized_to_date": 890.10, "days_covered": 26.0, "first_day": "2026-10-08",
            "last_day": "2026-11-02", "annualized_realized_savings": None, "expected_same_scope_annual_savings": 9363.51,
            "realization_ratio": None, "attribution_claim_status": "CLAIMABLE", "next_action": "None: rolled back",
            "as_of_utc": "2026-11-20T12:00:00Z"})

    def test_a_not_applied_lineage(self):
        (row,) = report.project_realization_workbook([Fixture.not_applied]).warehouses
        self.assertEqual({k: row[k] for k in ("warehouse_id", "lineage_basis", "applied_knob", "from_value", "applied_value",
                                              "applied_at_utc", "implementation_conformance", "validation_state", "realization_state",
                                              "maturity_state", "currency", "realized_to_date", "days_covered", "annualized_realized_savings",
                                              "expected_same_scope_annual_savings", "realization_ratio", "attribution_claim_status",
                                              "next_action")}, {
            "warehouse_id": "wh-demo-013", "lineage_basis": "NEWEST_ACTIVE_PACKAGE", "applied_knob": None, "from_value": None,
            "applied_value": None, "applied_at_utc": None, "implementation_conformance": "NOT_APPLIED", "validation_state": "NOT_READY",
            "realization_state": "NOT_APPLIED", "maturity_state": "NOT_APPLIED", "currency": None, "realized_to_date": None,
            "days_covered": None, "annualized_realized_savings": None, "expected_same_scope_annual_savings": None,
            "realization_ratio": None, "attribution_claim_status": "NOT_ATTRIBUTABLE", "next_action": "Apply the recommended change"})

    def test_a_not_applied_lineage_names_its_first_plan_step_when_the_run_gives_it(self):
        from realizer.analyzer.analysis import next_plan_step
        step = next_plan_step(Fixture.ge01)
        self.assertEqual(step, {"step": 1, "knob": "auto_stop_minutes", "from_value": 30, "to_value": 10})
        run = dataclasses.replace(Fixture.not_applied, next_steps={Fixture.ge01["recommendation_package_id"]: step})
        projection = report.project_realization_workbook([run])
        self.assertEqual(projection.warehouses[0]["next_action"], "Apply step 1: auto_stop_minutes 30 \u2192 10")
        self.assertEqual(projection.lineages[0]["next_action"], "Apply step 1: auto_stop_minutes 30 \u2192 10")
        # Only a lineage not yet applied names the step.
        applied = dataclasses.replace(Fixture.rolled_back, next_steps={Fixture.ge01["recommendation_package_id"]: step})
        self.assertEqual(report.project_realization_workbook([applied]).warehouses[0]["next_action"], "None: rolled back")

    def test_the_first_plan_step_skips_steps_a_person_cannot_apply_yet(self):
        from realizer.analyzer.analysis import next_plan_step
        self.assertEqual(next_plan_step(fixture_package("two-ready-steps")),
                         {"step": 1, "knob": "warehouse_size", "from_value": "MEDIUM", "to_value": "SMALL"})
        self.assertIsNone(next_plan_step(fixture_package("scope-migration-only")))

    def test_a_failed_warehouse_still_gets_a_row(self):
        rows = report.project_realization_workbook([CRASHED, REFUSED]).warehouses
        self.assertEqual([(r["warehouse_id"], r["run_status"], r["run_reason"], r["recommendation_package_id"], r["realization_package_id"],
                           r["lineage_basis"], r["realization_state"], r["other_lineages"], r["next_action"]) for r in rows], [
            ("wh-demo-020", "REFUSED", REFUSED.reason, REFUSED.recommendation_package_id, None, None, None, 0, "Fix the run: REFUSED"),
            ("wh-demo-021", "CRASHED", "KeyError", None, None, None, None, 0, "Fix the run: CRASHED"),
        ])

    def test_a_skipped_warehouse_gets_a_row_that_asks_for_nothing(self):
        skipped = WarehouseRun(workspace_id="1234567890123456", warehouse_id="wh-demo-022", status=report.RUN_STATUS_SKIPPED,
                               reason="the collector warehouse cannot measure itself")
        (row,) = report.project_realization_workbook([skipped]).warehouses
        self.assertEqual((row["run_status"], row["run_reason"], row["realization_state"], row["next_action"]),
                         ("SKIPPED", "the collector warehouse cannot measure itself", None, "None: collector warehouse, not measured"))
        self.assertIn(report.RUN_STATUS_SKIPPED, report.RUN_STATUSES)

    def test_a_failed_run_keeps_its_lineage_but_asks_to_fix_the_run(self):
        run = dataclasses.replace(Fixture.rolled_back, status=report.RUN_STATUS_REFUSED, reason="the ledger is locked")
        (row,) = report.project_realization_workbook([run]).warehouses
        self.assertEqual((row["realization_state"], row["next_action"]), ("ROLLED_BACK", "Fix the run: REFUSED"))


class WorkbookProjectionTests(unittest.TestCase):
    def test_rows_and_order(self):
        projection = report.project_realization_workbook([REFUSED, Fixture.rolled_back, Fixture.pair, CRASHED])
        self.assertEqual([(r["workspace_id"], r["warehouse_id"]) for r in projection.warehouses], [
            ("1234567890123456", "wh-demo-020"), ("1234567890123456", "wh-demo-021"), ("2234567890123456", "wh-demo-013"),
            ("2234567890123456", "wh-demo-039")])
        # One lineage row per realization package, oldest package first within a warehouse.
        self.assertEqual([r["recommendation_package_id"] for r in projection.lineages],
                         [Fixture.rolled_back.realizations[0].source_recommendation["recommendation_package_id"],
                          *sorted([Fixture.balanced_pkg, Fixture.aggressive_pkg])])
        # Every ledger event of every lineage, in ledger order within its lineage.
        self.assertEqual(len(projection.timeline), sum(len(v) for run in (Fixture.rolled_back, Fixture.pair) for v in run.ledgers.values()))
        for sheet, rows in ((report.SHEET_WAREHOUSES, projection.warehouses), (report.SHEET_LINEAGES, projection.lineages),
                            (report.SHEET_TIMELINE, projection.timeline)):
            with self.subTest(sheet=sheet):
                self.assertTrue(all(list(r) == keys(report.SHEET_COLUMNS[sheet][0]) for r in rows))
        self.assertEqual(projection.as_of_utc, "2026-11-20T12:00:00Z")

    def test_the_lineage_rows_of_the_profile_pair(self):
        lineages = report.project_realization_workbook([Fixture.pair]).lineages
        balanced, aggressive = by_rpkg(lineages, rc.make_rpkg_id(Fixture.balanced_pkg)), by_rpkg(lineages, rc.make_rpkg_id(Fixture.aggressive_pkg))
        pick = ("implementation_conformance", "recommended_value", "observed_applied_value", "attribution_claim_status",
                "claimed_by_package_id", "observed_realized_savings", "recommendation_attributed_savings", "realized_to_date",
                "realized_to_date_days_covered", "annualized_realized_savings", "currency", "regeneration_status", "next_action")
        self.assertEqual({k: balanced[k] for k in pick}, {
            "implementation_conformance": "EXACT", "recommended_value": 210, "observed_applied_value": 210,
            "attribution_claim_status": "CLAIMABLE", "claimed_by_package_id": None, "observed_realized_savings": 1069.73,
            "recommendation_attributed_savings": 1069.73, "realized_to_date": 1069.73, "realized_to_date_days_covered": 28.0,
            "annualized_realized_savings": 13944.69, "currency": "USD", "regeneration_status": "NOT_REQUIRED", "next_action": "Keep monitoring"})
        self.assertEqual({k: aggressive[k] for k in pick}, {
            "implementation_conformance": "MODIFIED", "recommended_value": 15, "observed_applied_value": 210,
            "attribution_claim_status": "OBSERVABLE_ONLY", "claimed_by_package_id": Fixture.balanced_pkg, "observed_realized_savings": 1069.73,
            "recommendation_attributed_savings": 0.0, "realized_to_date": 1069.73, "realized_to_date_days_covered": 28.0,
            "annualized_realized_savings": 13944.69, "currency": "USD", "regeneration_status": "REQUESTED", "next_action": "Keep monitoring"})

    def test_the_lineage_projection_does_not_recompute(self):
        (lineage,) = report.project_realization_workbook([Fixture.rolled_back]).lineages
        source = Fixture.rolled_back.realizations[0]
        self.assertEqual(lineage["realized_to_date"], float(source.economics["realized_savings_since_application"]))
        self.assertEqual(json.loads(lineage["economics_json"]), source.economics)
        self.assertEqual({k: v for k, v in lineage.items() if k in report.project_summary(source)}, report.project_summary(source))

    def test_a_duplicate_warehouse_is_refused(self):
        with self.assertRaises(ValueError):
            report.project_realization_workbook([Fixture.pair, Fixture.pair])


class RenderedWorkbookTests(unittest.TestCase):
    """The file, opened back with openpyxl."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.runs = [Fixture.pair, Fixture.rolled_back, REFUSED]
        cls.path = write_realization_workbook(Path(cls.tmp.name) / "nested" / "realization.xlsx", cls.runs)
        cls.book = load_workbook(cls.path)
        cls.projection = report.project_realization_workbook(cls.runs)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def table(self, sheet: str) -> Tuple[List[str], List[Dict]]:
        ws = self.book[sheet]
        rows = list(ws.iter_rows(min_row=report_header_row(), values_only=True))
        header, body = list(rows[0]), rows[1:]
        return header, [dict(zip(header, r)) for r in body]

    def test_exactly_three_visible_sheets(self):
        self.assertEqual(self.book.sheetnames, list(report.SHEET_NAMES))
        self.assertTrue(all(ws.sheet_state == "visible" for ws in self.book.worksheets))
        self.assertFalse(self.path.with_suffix(".xlsm").exists())
        self.assertIsNone(getattr(self.book, "vba_archive", None))

    def test_headers_in_contract_order_and_one_row_per_record(self):
        expected_rows = {report.SHEET_WAREHOUSES: 3, report.SHEET_LINEAGES: 3, report.SHEET_TIMELINE: len(self.projection.timeline)}
        for sheet, (columns, groups) in report.SHEET_COLUMNS.items():
            with self.subTest(sheet=sheet):
                header, body = self.table(sheet)
                self.assertEqual(header, [c.header for c in columns])
                self.assertEqual(len(body), expected_rows[sheet])
                ws = self.book[sheet]
                self.assertEqual(ws.freeze_panes[1:], str(report_header_row() + 1))
                group_titles = [ws.cell(row=report_header_row() - 1, column=keys(columns).index(g.start_key) + 1).value for g in groups]
                self.assertEqual(group_titles, [g.title for g in groups])

    def test_cells_match_the_projection(self):
        for sheet, rows in ((report.SHEET_WAREHOUSES, self.projection.warehouses), (report.SHEET_LINEAGES, self.projection.lineages),
                            (report.SHEET_TIMELINE, self.projection.timeline)):
            columns = report.SHEET_COLUMNS[sheet][0]
            _, body = self.table(sheet)
            with self.subTest(sheet=sheet):
                self.assertEqual([[r[c.header] for c in columns] for r in body], [[row[c.key] for c in columns] for row in rows])

    def test_money_is_numeric_and_signed_ids_are_text_and_unknown_is_empty(self):
        header, body = self.table(report.SHEET_WAREHOUSES)
        rows = {r["Warehouse ID"]: r for r in body}
        self.assertEqual(rows["wh-demo-013"]["Realized to date"], 890.10)
        self.assertIsNone(rows["wh-demo-013"]["Annualized realized"])  # unknown: empty, never zero
        self.assertIsNone(rows["wh-demo-020"]["Realized to date"])
        self.assertEqual(rows["wh-demo-039"]["Workspace ID"], "2234567890123456")
        ws = self.book[report.SHEET_WAREHOUSES]
        for spec_index, spec in enumerate(report.WAREHOUSE_COLUMNS, start=1):
            for row in range(report_header_row() + 1, ws.max_row + 1):
                cell = ws.cell(row=row, column=spec_index)
                if cell.value is None:
                    continue
                if spec.kind == report.KIND_ID:
                    self.assertEqual((cell.data_type, cell.number_format), ("s", "@"), spec.key)
                if spec.kind == report.KIND_MONEY:
                    self.assertIsInstance(cell.value, float, spec.key)
                    self.assertTrue(cell.number_format.startswith("#,##0.00;"), spec.key)
        # A negative realization stays negative: render a signed copy of the rolled-back lineage.
        source = Fixture.rolled_back.realizations[0]
        negative = dataclasses.replace(source, economics={**source.economics, "realized_savings_since_application": "-12.34"})
        with tempfile.TemporaryDirectory() as tmp:
            path = write_realization_workbook(Path(tmp) / "w.xlsx", [dataclasses.replace(Fixture.rolled_back, realizations=(negative,))])
            ws = load_workbook(path)[report.SHEET_WAREHOUSES]
            column = keys(report.WAREHOUSE_COLUMNS).index("realized_to_date") + 1
            self.assertEqual(ws.cell(row=report_header_row() + 1, column=column).value, -12.34)

    def test_no_formulas_anywhere(self):
        for ws in self.book.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", f"{ws.title}!{cell.coordinate}")
            self.assertEqual(len(ws.conditional_formatting), 0)
        self.assertEqual(len(self.book.defined_names), 0)

    def test_identical_inputs_give_identical_content_and_properties(self):
        with tempfile.TemporaryDirectory() as tmp:
            again = load_workbook(write_realization_workbook(Path(tmp) / "again.xlsx", self.runs))
        for sheet in report.SHEET_NAMES:
            self.assertEqual([[c.value for c in r] for r in again[sheet].iter_rows()], [[c.value for c in r] for r in self.book[sheet].iter_rows()])
        self.assertEqual(again.properties.created, self.book.properties.created)
        self.assertEqual(self.book.properties.created.isoformat(), "2026-11-20T12:00:00")
        self.assertEqual(self.book.properties.modified, self.book.properties.created)

    def test_an_empty_fleet_still_writes_three_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            book = load_workbook(write_realization_workbook(Path(tmp) / "empty.xlsx", []))
        self.assertEqual(book.sheetnames, list(report.SHEET_NAMES))


def report_header_row() -> int:
    from realizer.reporter import workbook
    return workbook.HEADER_ROW


if __name__ == "__main__":
    unittest.main()
